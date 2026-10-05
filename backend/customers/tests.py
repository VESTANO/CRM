import io
import json
from datetime import date, time, timedelta
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from .models import AssignedTask, CustomerRecord, Dataset, DatasetColumn, ManualClient, Reminder, SalesTarget, Visit


class ExcelInspectionTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="crmuser",
            password="test-password-123",
        )
        self.client.force_login(self.user)

    def test_import_page_inspects_before_confirming_dataset_import(self):
        response = self.client.post(
            reverse("customers:import_excel"),
            {
                "dataset_name": "January Leads",
                "excel_file": SimpleUploadedFile(
                    "customers.xlsx",
                    build_workbook_bytes(
                        "Leads",
                        ["Customer", "Contact", "Amount"],
                        [
                            ["Asha Rao", "asha.rao@example.com", 1200],
                            ["Ravi Shah", "9876543210", 900],
                        ],
                    ),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Workbook Summary")
        self.assertContains(response, "Confirm Import")
        self.assertContains(response, "customers.xlsx")
        self.assertContains(response, "Leads")
        self.assertContains(response, "Customer")
        self.assertContains(response, "a***o@example.com")
        self.assertEqual(Dataset.objects.count(), 0)

    def test_confirm_import_creates_isolated_dataset_with_columns_and_rows(self):
        dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="customers.xlsx",
            headers=["Customer", "Contact", "Amount"],
            rows=[
                ["Asha Rao", "asha.rao@example.com", 1200],
                ["Ravi Shah", "9876543210", 900],
            ],
        )

        self.assertEqual(Dataset.objects.count(), 1)
        dataset = Dataset.objects.get(name="January Leads")
        self.assertEqual(dataset.original_filename, "customers.xlsx")
        self.assertEqual(dataset.records.count(), 2)
        self.assertEqual(dataset.columns.count(), 3)
        self.assertEqual(
            list(dataset.columns.values_list("original_name", flat=True)),
            ["Customer", "Contact", "Amount"],
        )

        first_record = dataset.records.order_by("row_number").first()
        self.assertEqual(first_record.original_data["customer"], "Asha Rao")
        self.assertEqual(first_record.original_data["contact"], "asha.rao@example.com")
        self.assertEqual(first_record.response, CustomerRecord.RESPONSE_NO_RESPONSE)
        self.assertEqual(first_record.note, "")

        detail_response = self.client.get(reverse("customers:dataset_detail", args=[dataset.id]))
        self.assertContains(detail_response, "a***o@example.com")
        self.assertContains(detail_response, "XXXXXXXX10")
        self.assertContains(detail_response, "Response")
        self.assertContains(detail_response, "Note")
        self.assertContains(detail_response, "View/Edit")

    def test_import_page_rejects_non_xlsx_files(self):
        response = self.client.post(
            reverse("customers:import_excel"),
            {
                "dataset_name": "January Leads",
                "excel_file": SimpleUploadedFile(
                    "customers.csv",
                    b"name,email\nAsha,asha@example.com\n",
                    content_type="text/csv",
                )
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Please upload a valid .xlsx Excel file.")
        self.assertEqual(Dataset.objects.count(), 0)

    def test_multiple_imports_create_independent_datasets(self):
        first_dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name", "Phone Number"],
            rows=[["Asha Rao", "1111111111"]],
        )
        second_dataset = import_dataset(
            self.client,
            dataset_name="February Leads",
            filename="february.xlsx",
            headers=["Company", "Product", "Location"],
            rows=[
                ["Northwind", "Starter", "Pune"],
                ["Contoso", "Pro", "Delhi"],
            ],
        )

        self.assertEqual(Dataset.objects.count(), 2)
        self.assertEqual(first_dataset.records.count(), 1)
        self.assertEqual(second_dataset.records.count(), 2)
        self.assertEqual(
            list(first_dataset.columns.values_list("original_name", flat=True)),
            ["Customer Name", "Phone Number"],
        )
        self.assertEqual(
            list(second_dataset.columns.values_list("original_name", flat=True)),
            ["Company", "Product", "Location"],
        )
        self.assertFalse(
            first_dataset.records.filter(original_data__has_key="company").exists()
        )
        self.assertFalse(
            second_dataset.records.filter(original_data__has_key="phone_number").exists()
        )

    def test_empty_rows_are_skipped_without_creating_invalid_records(self):
        response = self.client.post(
            reverse("customers:import_excel"),
            {
                "dataset_name": "Empty Row Leads",
                "excel_file": SimpleUploadedFile(
                    "empty_rows.xlsx",
                    build_workbook_bytes(
                        "Leads",
                        ["Customer Name", "Phone Number"],
                        [["", ""], [None, None]],
                    ),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )
        token = response.context["pending_import_token"]
        confirm_response = self.client.post(
            reverse("customers:import_excel"),
            {
                "action": "confirm_import",
                "pending_import_token": token,
            },
        )

        dataset = Dataset.objects.get(name="Empty Row Leads")
        self.assertRedirects(confirm_response, reverse("customers:dataset_detail", args=[dataset.id]))
        self.assertEqual(dataset.records.count(), 0)
        detail_response = self.client.get(reverse("customers:dataset_detail", args=[dataset.id]))
        self.assertContains(detail_response, "No customer records were found in this file.")

    def test_duplicate_column_names_are_preserved_with_unique_keys(self):
        dataset = import_dataset(
            self.client,
            dataset_name="Duplicate Header Leads",
            filename="duplicate_headers.xlsx",
            headers=["Phone", "Phone"],
            rows=[["1111111111", "2222222222"]],
        )

        self.assertEqual(
            list(DatasetColumn.objects.filter(dataset=dataset).values_list("original_name", flat=True)),
            ["Phone", "Phone"],
        )
        self.assertEqual(
            list(DatasetColumn.objects.filter(dataset=dataset).values_list("field_key", flat=True)),
            ["phone", "phone_2"],
        )

    def test_dataset_detail_search_is_dataset_scoped(self):
        first_dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name", "City"],
            rows=[["Rahul Nair", "Kozhikode"], ["Meera Nair", "Kochi"]],
        )
        second_dataset = import_dataset(
            self.client,
            dataset_name="February Leads",
            filename="february.xlsx",
            headers=["Customer Name", "City"],
            rows=[["Rahul Sharma", "Delhi"]],
        )

        response = self.client.get(
            reverse("customers:dataset_detail", args=[first_dataset.id]),
            {"q": "Rahul"},
        )

        self.assertContains(response, "Rahul Nair")
        self.assertNotContains(response, "Meera Nair")
        self.assertNotContains(response, "Rahul Sharma")
        self.assertEqual(first_dataset.records.count(), 2)
        self.assertEqual(second_dataset.records.count(), 1)

    def test_response_filter_and_search_work_together(self):
        dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name", "City"],
            rows=[["Rahul Nair", "Kozhikode"], ["Rahul Menon", "Kochi"]],
        )
        interested = dataset.records.get(original_data__customer_name="Rahul Nair")
        interested.response = CustomerRecord.RESPONSE_INTERESTED
        interested.save()

        response = self.client.get(
            reverse("customers:dataset_detail", args=[dataset.id]),
            {
                "q": "Rahul",
                "response": CustomerRecord.RESPONSE_INTERESTED,
            },
        )

        self.assertContains(response, "Rahul Nair")
        self.assertNotContains(response, "Rahul Menon")
        self.assertContains(response, "Showing 1 of 2 customers.")

    def test_customer_detail_and_edit_update_only_crm_fields(self):
        dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name", "City"],
            rows=[["Rahul Nair", "Kozhikode"]],
        )
        record = dataset.records.get()
        original_data = record.original_data.copy()

        detail_response = self.client.get(
            reverse("customers:customer_detail", args=[dataset.id, record.id])
        )
        self.assertContains(detail_response, "Rahul Nair")
        self.assertContains(detail_response, "No Response")

        edit_response = self.client.post(
            reverse("customers:customer_edit", args=[dataset.id, record.id]),
            {
                "response": CustomerRecord.RESPONSE_CALLBACK,
                "note": "Call again tomorrow regarding pricing.",
            },
        )

        self.assertRedirects(
            edit_response,
            reverse("customers:customer_detail", args=[dataset.id, record.id]),
        )
        record.refresh_from_db()
        self.assertEqual(record.response, CustomerRecord.RESPONSE_CALLBACK)
        self.assertEqual(record.note, "Call again tomorrow regarding pricing.")
        self.assertEqual(record.original_data, original_data)

    def test_delete_customer_is_post_only_and_dataset_scoped(self):
        first_dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name"],
            rows=[["Asha Rao"]],
        )
        second_dataset = import_dataset(
            self.client,
            dataset_name="February Leads",
            filename="february.xlsx",
            headers=["Customer Name"],
            rows=[["Asha Rao"]],
        )
        first_record = first_dataset.records.get()
        second_record = second_dataset.records.get()

        get_response = self.client.get(
            reverse("customers:customer_delete", args=[first_dataset.id, first_record.id])
        )
        self.assertEqual(get_response.status_code, 200)
        self.assertEqual(first_dataset.records.count(), 1)

        delete_response = self.client.post(
            reverse("customers:customer_delete", args=[first_dataset.id, first_record.id])
        )

        self.assertRedirects(
            delete_response,
            reverse("customers:dataset_detail", args=[first_dataset.id]),
        )
        self.assertFalse(CustomerRecord.objects.filter(id=first_record.id).exists())
        self.assertTrue(CustomerRecord.objects.filter(id=second_record.id).exists())

    def test_wrong_dataset_customer_access_returns_404(self):
        first_dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name"],
            rows=[["Asha Rao"]],
        )
        second_dataset = import_dataset(
            self.client,
            dataset_name="February Leads",
            filename="february.xlsx",
            headers=["Customer Name"],
            rows=[["Meera Nair"]],
        )
        second_record = second_dataset.records.get()

        response = self.client.get(
            reverse("customers:customer_detail", args=[first_dataset.id, second_record.id])
        )

        self.assertEqual(response.status_code, 404)

    def test_dataset_detail_paginates_customers(self):
        rows = [[f"Customer {index}"] for index in range(1, 26)]
        dataset = import_dataset(
            self.client,
            dataset_name="Large Leads",
            filename="large.xlsx",
            headers=["Customer Name"],
            rows=rows,
        )

        first_page = self.client.get(reverse("customers:dataset_detail", args=[dataset.id]))
        second_page = self.client.get(
            reverse("customers:dataset_detail", args=[dataset.id]),
            {"page": 2},
        )

        self.assertContains(first_page, "Customer 1")
        self.assertContains(first_page, "Customer 20")
        self.assertNotContains(first_page, "Customer 21")
        self.assertContains(second_page, "Customer 21")
        self.assertContains(second_page, "Customer 25")

    def test_export_entire_dataset_preserves_columns_and_adds_crm_fields(self):
        dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name", "Phone Number", "Email"],
            rows=[["Rahul Nair", "1111111111", "rahul@example.com"]],
        )
        record = dataset.records.get()
        record.response = CustomerRecord.RESPONSE_INTERESTED
        record.note = "Call tomorrow"
        record.save()

        response = self.client.get(reverse("customers:dataset_export", args=[dataset.id]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("january_leads.xlsx", response["Content-Disposition"])
        rows = read_export_rows(response.content)
        self.assertEqual(
            rows[0],
            ["Customer Name", "Phone Number", "Email", "Response", "Note"],
        )
        self.assertEqual(
            rows[1],
            ["Rahul Nair", "1111111111", "rahul@example.com", "Interested", "Call tomorrow"],
        )

    def test_export_search_and_response_filter_results_only(self):
        dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name", "City"],
            rows=[
                ["Rahul Nair", "Kozhikode"],
                ["Rahul Menon", "Kochi"],
                ["Meera Nair", "Kozhikode"],
            ],
        )
        interested = dataset.records.get(original_data__customer_name="Rahul Nair")
        interested.response = CustomerRecord.RESPONSE_INTERESTED
        interested.save()

        response = self.client.get(
            reverse("customers:dataset_export", args=[dataset.id]),
            {
                "q": "Rahul",
                "response": CustomerRecord.RESPONSE_INTERESTED,
            },
        )

        rows = read_export_rows(response.content)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], "Rahul Nair")
        self.assertNotIn("Rahul Menon", [row[0] for row in rows[1:]])
        self.assertNotIn("Meera Nair", [row[0] for row in rows[1:]])

    def test_dataset_can_be_renamed_without_changing_records_or_columns(self):
        dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name"],
            rows=[["Rahul Nair"]],
        )

        response = self.client.post(
            reverse("customers:dataset_rename", args=[dataset.id]),
            {"name": "January Sales Leads"},
        )

        self.assertRedirects(response, reverse("customers:dataset_detail", args=[dataset.id]))
        dataset.refresh_from_db()
        self.assertEqual(dataset.name, "January Sales Leads")
        self.assertEqual(dataset.original_filename, "january.xlsx")
        self.assertEqual(dataset.columns.get().original_name, "Customer Name")
        self.assertEqual(dataset.records.get().original_data["customer_name"], "Rahul Nair")

    def test_dataset_delete_removes_only_that_dataset_and_its_records(self):
        first_dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name"],
            rows=[["Rahul Nair"], ["Meera Nair"]],
        )
        second_dataset = import_dataset(
            self.client,
            dataset_name="February Leads",
            filename="february.xlsx",
            headers=["Customer Name"],
            rows=[["Asha Rao"]],
        )

        response = self.client.post(reverse("customers:dataset_delete", args=[first_dataset.id]))

        self.assertRedirects(response, reverse("customers:dataset_list"))
        self.assertFalse(Dataset.objects.filter(id=first_dataset.id).exists())
        self.assertTrue(Dataset.objects.filter(id=second_dataset.id).exists())
        self.assertEqual(CustomerRecord.objects.filter(dataset=second_dataset).count(), 1)
        self.assertEqual(CustomerRecord.objects.count(), 1)

    def test_dashboard_counts_and_recent_datasets_are_correct(self):
        import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name"],
            rows=[["Rahul Nair"], ["Meera Nair"]],
        )
        import_dataset(
            self.client,
            dataset_name="February Leads",
            filename="february.xlsx",
            headers=["Customer Name"],
            rows=[["Asha Rao"]],
        )

        response = self.client.get(reverse("customers:dashboard"))

        self.assertContains(response, "Datasets")
        self.assertContains(response, "2")
        self.assertContains(response, "Total Customers")
        self.assertContains(response, "3")
        self.assertContains(response, "January Leads")
        self.assertContains(response, "February Leads")

    def test_dataset_response_statistics_are_dataset_specific(self):
        first_dataset = import_dataset(
            self.client,
            dataset_name="January Leads",
            filename="january.xlsx",
            headers=["Customer Name"],
            rows=[["Rahul Nair"], ["Meera Nair"]],
        )
        second_dataset = import_dataset(
            self.client,
            dataset_name="February Leads",
            filename="february.xlsx",
            headers=["Customer Name"],
            rows=[["Asha Rao"]],
        )
        first_record = first_dataset.records.first()
        first_record.response = CustomerRecord.RESPONSE_CONVERTED
        first_record.save()
        second_record = second_dataset.records.first()
        second_record.response = CustomerRecord.RESPONSE_CONVERTED
        second_record.save()

        response = self.client.get(reverse("customers:dataset_detail", args=[first_dataset.id]))

        self.assertContains(response, "Converted")
        self.assertContains(response, '<div class="fs-5">1</div>', html=True)


class AuthenticationTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="crmuser",
            password="test-password-123",
        )

    def test_login_page_is_accessible_without_authentication(self):
        response = self.client.get(reverse("login"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Customer CRM")
        self.assertContains(response, "Username")
        self.assertContains(response, "Password")

    def test_dashboard_requires_authentication(self):
        response = self.client.get(reverse("customers:dashboard"))

        self.assertRedirects(
            response,
            f"{reverse('login')}?next={reverse('customers:dashboard')}",
        )

    def test_customer_lists_requires_authentication(self):
        response = self.client.get(reverse("customers:dataset_list"))

        self.assertRedirects(
            response,
            f"{reverse('login')}?next={reverse('customers:dataset_list')}",
        )

    def test_import_page_requires_authentication(self):
        response = self.client.get(reverse("customers:import_excel"))

        self.assertRedirects(
            response,
            f"{reverse('login')}?next={reverse('customers:import_excel')}",
        )

    def test_invalid_login_is_rejected(self):
        response = self.client.post(
            reverse("login"),
            {
                "username": "crmuser",
                "password": "wrong-password",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Invalid username or password.")
        self.assertNotIn("_auth_user_id", self.client.session)

    def test_valid_login_succeeds_and_redirects_to_dashboard(self):
        response = self.client.post(
            reverse("login"),
            {
                "username": "crmuser",
                "password": "test-password-123",
            },
        )

        self.assertRedirects(response, reverse("customers:dashboard"))
        self.assertEqual(
            self.client.session["_auth_user_id"],
            str(self.user.id),
        )

    def test_authenticated_user_can_access_crm_pages(self):
        self.client.force_login(self.user)

        dashboard_response = self.client.get(reverse("customers:dashboard"))
        dataset_list_response = self.client.get(reverse("customers:dataset_list"))
        import_response = self.client.get(reverse("customers:import_excel"))

        self.assertEqual(dashboard_response.status_code, 200)
        self.assertEqual(dataset_list_response.status_code, 200)
        self.assertEqual(import_response.status_code, 200)

    def test_logout_works_and_protected_pages_require_login_again(self):
        self.client.force_login(self.user)

        logout_response = self.client.post(reverse("logout"))
        self.assertRedirects(logout_response, reverse("login"))
        self.assertNotIn("_auth_user_id", self.client.session)

        dashboard_response = self.client.get(reverse("customers:dashboard"))
        self.assertRedirects(
            dashboard_response,
            f"{reverse('login')}?next={reverse('customers:dashboard')}",
        )


class OwnershipAccessTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(
            username="admin",
            password="test-password-123",
            is_staff=True,
        )
        self.user_one = User.objects.create_user(
            username="userone",
            password="test-password-123",
        )
        self.user_two = User.objects.create_user(
            username="usertwo",
            password="test-password-123",
        )

        self.client.force_login(self.user_one)
        self.user_one_dataset = import_dataset(
            self.client,
            dataset_name="User One Leads",
            filename="one.xlsx",
            headers=["Customer Name"],
            rows=[["Rahul Nair"], ["Meera Nair"]],
        )
        user_one_interested = self.user_one_dataset.records.first()
        user_one_interested.response = CustomerRecord.RESPONSE_INTERESTED
        user_one_interested.save()

        self.client.force_login(self.user_two)
        self.user_two_dataset = import_dataset(
            self.client,
            dataset_name="User Two Leads",
            filename="two.xlsx",
            headers=["Customer Name"],
            rows=[["Asha Rao"]],
        )
        self.user_two_record = self.user_two_dataset.records.get()

    def test_admin_can_see_all_datasets(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse("customers:dataset_list"))

        self.assertContains(response, "User One Leads")
        self.assertContains(response, "User Two Leads")
        self.assertContains(response, "userone")
        self.assertContains(response, "usertwo")

    def test_admin_can_see_all_customers(self):
        self.client.force_login(self.admin)

        first_response = self.client.get(
            reverse("customers:dataset_detail", args=[self.user_one_dataset.id])
        )
        second_response = self.client.get(
            reverse("customers:dataset_detail", args=[self.user_two_dataset.id])
        )

        self.assertContains(first_response, "Rahul Nair")
        self.assertContains(second_response, "Asha Rao")

    def test_user_can_see_own_dataset(self):
        self.client.force_login(self.user_one)

        response = self.client.get(
            reverse("customers:dataset_detail", args=[self.user_one_dataset.id])
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "User One Leads")
        self.assertContains(response, "Rahul Nair")

    def test_user_cannot_see_another_users_dataset(self):
        self.client.force_login(self.user_one)

        response = self.client.get(
            reverse("customers:dataset_detail", args=[self.user_two_dataset.id])
        )

        self.assertEqual(response.status_code, 404)

    def test_user_cannot_access_another_users_customer_detail(self):
        self.client.force_login(self.user_one)

        response = self.client.get(
            reverse(
                "customers:customer_detail",
                args=[self.user_two_dataset.id, self.user_two_record.id],
            )
        )

        self.assertEqual(response.status_code, 404)

    def test_user_cannot_edit_another_users_customer(self):
        self.client.force_login(self.user_one)

        response = self.client.post(
            reverse(
                "customers:customer_edit",
                args=[self.user_two_dataset.id, self.user_two_record.id],
            ),
            {
                "response": CustomerRecord.RESPONSE_CONVERTED,
                "note": "Should not save",
            },
        )

        self.assertEqual(response.status_code, 404)
        self.user_two_record.refresh_from_db()
        self.assertNotEqual(self.user_two_record.response, CustomerRecord.RESPONSE_CONVERTED)
        self.assertEqual(self.user_two_record.note, "")

    def test_user_cannot_delete_another_users_customer(self):
        self.client.force_login(self.user_one)

        response = self.client.post(
            reverse(
                "customers:customer_delete",
                args=[self.user_two_dataset.id, self.user_two_record.id],
            )
        )

        self.assertEqual(response.status_code, 404)
        self.assertTrue(CustomerRecord.objects.filter(id=self.user_two_record.id).exists())

    def test_user_cannot_export_another_users_dataset(self):
        self.client.force_login(self.user_one)

        response = self.client.get(
            reverse("customers:dataset_export", args=[self.user_two_dataset.id])
        )

        self.assertEqual(response.status_code, 404)

    def test_user_dashboard_only_shows_own_statistics(self):
        self.client.force_login(self.user_one)

        response = self.client.get(reverse("customers:dashboard"))

        self.assertContains(response, "User One Leads")
        self.assertNotContains(response, "User Two Leads")
        self.assertContains(response, "Total Customers")
        self.assertContains(response, "2")
        self.assertContains(response, "Interested")

    def test_importing_dataset_assigns_logged_in_user_as_owner(self):
        self.client.force_login(self.user_one)

        dataset = import_dataset(
            self.client,
            dataset_name="Owned Import",
            filename="owned.xlsx",
            headers=["Customer Name"],
            rows=[["Nila Shah"]],
        )

        self.assertEqual(dataset.owner, self.user_one)


class CustomAdminPanelTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(
            username="admin",
            password="test-password-123",
            email="admin@example.com",
            is_staff=True,
        )
        self.user = User.objects.create_user(
            username="crmuser",
            password="test-password-123",
            email="crmuser@example.com",
        )
        self.other_user = User.objects.create_user(
            username="otheruser",
            password="test-password-123",
        )

        self.client.force_login(self.user)
        self.user_dataset = import_dataset(
            self.client,
            dataset_name="User Owned Leads",
            filename="owned.xlsx",
            headers=["Customer Name"],
            rows=[["Rahul Nair"], ["Meera Nair"]],
        )

        self.client.force_login(self.other_user)
        self.other_dataset = import_dataset(
            self.client,
            dataset_name="Other User Leads",
            filename="other.xlsx",
            headers=["Customer Name"],
            rows=[["Asha Rao"]],
        )

    def test_admin_can_access_admin_panel(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse("customers:admin_panel"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Admin Panel")
        self.assertContains(response, "Manage users, datasets, and CRM activity.")

    def test_normal_user_cannot_access_admin_panel(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("customers:admin_panel"))

        self.assertEqual(response.status_code, 403)

    def test_admin_panel_navigation_points_to_custom_admin_panel(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse("customers:dashboard"))

        self.assertContains(response, reverse("customers:admin_panel"))
        self.assertNotContains(response, 'href="/admin/"')

    def test_django_admin_remains_available_separately(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse("admin:index"))

        self.assertEqual(response.status_code, 200)

    def test_admin_can_view_users(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse("customers:admin_panel"))

        self.assertContains(response, "crmuser")
        self.assertContains(response, "otheruser")
        self.assertContains(response, "Admin")
        self.assertContains(response, "User")

    def test_admin_can_create_normal_user(self):
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse("customers:admin_user_add"),
            {
                "username": "newcrmuser",
                "email": "newcrmuser@example.com",
                "password1": "Strong-pass-12345",
                "password2": "Strong-pass-12345",
                "is_active": "on",
            },
        )

        new_user = get_user_model().objects.get(username="newcrmuser")
        self.assertRedirects(
            response,
            reverse("customers:admin_user_detail", args=[new_user.id]),
        )
        self.assertFalse(new_user.is_staff)
        self.assertFalse(new_user.is_superuser)
        self.assertTrue(new_user.check_password("Strong-pass-12345"))

    def test_admin_can_edit_user(self):
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse("customers:admin_user_edit", args=[self.user.id]),
            {
                "username": "crmuser-renamed",
                "email": "renamed@example.com",
                "is_active": "on",
            },
        )

        self.assertRedirects(
            response,
            reverse("customers:admin_user_detail", args=[self.user.id]),
        )
        self.user.refresh_from_db()
        self.assertEqual(self.user.username, "crmuser-renamed")
        self.assertEqual(self.user.email, "renamed@example.com")
        self.assertTrue(self.user.is_active)

    def test_admin_can_deactivate_user(self):
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse("customers:admin_user_toggle_active", args=[self.user.id])
        )

        self.assertRedirects(
            response,
            reverse("customers:admin_user_detail", args=[self.user.id]),
        )
        self.user.refresh_from_db()
        self.assertFalse(self.user.is_active)

    def test_admin_cannot_deactivate_self(self):
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse("customers:admin_user_toggle_active", args=[self.admin.id])
        )

        self.assertRedirects(
            response,
            reverse("customers:admin_user_detail", args=[self.admin.id]),
        )
        self.admin.refresh_from_db()
        self.assertTrue(self.admin.is_active)

    def test_admin_can_reset_normal_user_password(self):
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse("customers:admin_user_password", args=[self.user.id]),
            {
                "password1": "Reset-pass-12345",
                "password2": "Reset-pass-12345",
            },
        )

        self.assertRedirects(
            response,
            reverse("customers:admin_user_detail", args=[self.user.id]),
        )
        self.user.refresh_from_db()
        self.assertTrue(self.user.check_password("Reset-pass-12345"))

    def test_admin_can_view_users_datasets(self):
        self.client.force_login(self.admin)

        response = self.client.get(
            reverse("customers:admin_user_datasets", args=[self.user.id])
        )

        self.assertContains(response, "User Owned Leads")
        self.assertNotContains(response, "Other User Leads")
        self.assertContains(response, "Open dataset")

    def test_normal_user_cannot_access_user_management_urls(self):
        self.client.force_login(self.user)
        urls = [
            reverse("customers:admin_user_add"),
            reverse("customers:admin_user_detail", args=[self.user.id]),
            reverse("customers:admin_user_edit", args=[self.user.id]),
            reverse("customers:admin_user_password", args=[self.user.id]),
            reverse("customers:admin_user_datasets", args=[self.user.id]),
        ]

        for url in urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 403)

        toggle_response = self.client.post(
            reverse("customers:admin_user_toggle_active", args=[self.user.id])
        )
        self.assertEqual(toggle_response.status_code, 403)

    def test_admin_dashboard_statistics_are_correct(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse("customers:admin_panel"))

        self.assertContains(response, "Total Users")
        self.assertContains(response, "Active Users")
        self.assertContains(response, "Total Datasets")
        self.assertContains(response, "Total Customers")
        self.assertContains(response, "3")
        self.assertContains(response, "2")


class SalesVisitTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(
            username="salesadmin",
            password="test-password-123",
            is_staff=True,
        )
        self.user_one = User.objects.create_user(
            username="salesone",
            password="test-password-123",
        )
        self.user_two = User.objects.create_user(
            username="salestwo",
            password="test-password-123",
        )

        self.client.force_login(self.user_one)
        self.user_one_dataset = import_dataset(
            self.client,
            dataset_name="User One Customers",
            filename="one.xlsx",
            headers=["Customer Name", "City"],
            rows=[["Rahul Nair", "Kozhikode"], ["Meera Nair", "Kochi"]],
        )
        self.user_one_customer = self.user_one_dataset.records.first()

        self.client.force_login(self.user_two)
        self.user_two_dataset = import_dataset(
            self.client,
            dataset_name="User Two Customers",
            filename="two.xlsx",
            headers=["Customer Name", "City"],
            rows=[["Asha Rao", "Delhi"]],
        )
        self.user_two_customer = self.user_two_dataset.records.first()
        self.user_two_visit = Visit.objects.create(
            customer=self.user_two_customer,
            visit_date=date(2026, 9, 30),
            meeting_time=time(10, 30),
            location="Delhi Office",
            expense=Decimal("500.00"),
            status=Visit.STATUS_HEALTHY,
            note="Existing visit",
        )

    def test_dataset_client_selection_is_filtered_by_user_and_dataset(self):
        self.client.force_login(self.user_one)

        response = self.client.get(
            reverse("customers:visit_add"),
            {
                "dataset": self.user_one_dataset.id,
                "customer_q": "Rahul",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Rahul Nair")
        self.assertNotContains(response, "Meera Nair")
        self.assertNotContains(response, "Asha Rao")

    def test_sales_clients_endpoint_loads_dataset_customers(self):
        self.client.force_login(self.user_one)

        response = self.client.get(
            reverse("customers:sales_clients"),
            {"dataset": self.user_one_dataset.id},
        )

        self.assertEqual(response.status_code, 200)
        client_labels = [client["label"] for client in response.json()["clients"]]
        self.assertEqual(client_labels, ["Rahul Nair", "Meera Nair"])

    def test_sales_clients_endpoint_filters_customers_by_search(self):
        self.client.force_login(self.user_one)

        response = self.client.get(
            reverse("customers:sales_clients"),
            {
                "dataset": self.user_one_dataset.id,
                "q": "Kochi",
            },
        )

        self.assertEqual(response.status_code, 200)
        client_labels = [client["label"] for client in response.json()["clients"]]
        self.assertEqual(client_labels, ["Meera Nair"])

    def test_sales_clients_endpoint_enforces_dataset_ownership(self):
        self.client.force_login(self.user_one)

        response = self.client.get(
            reverse("customers:sales_clients"),
            {"dataset": self.user_two_dataset.id},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["clients"], [])

    def test_admin_sales_clients_endpoint_can_load_any_dataset(self):
        self.client.force_login(self.admin)

        response = self.client.get(
            reverse("customers:sales_clients"),
            {"dataset": self.user_two_dataset.id},
        )

        self.assertEqual(response.status_code, 200)
        client_labels = [client["label"] for client in response.json()["clients"]]
        self.assertEqual(client_labels, ["Asha Rao"])

    def test_create_visit_assigns_customer_in_owned_dataset(self):
        self.client.force_login(self.user_one)

        response = self.client.post(
            reverse("customers:visit_add"),
            {
                "dataset": self.user_one_dataset.id,
                "customer": self.user_one_customer.id,
                "customer_search": "",
                "visit_date": "2026-09-30",
                "meeting_time": "14:45",
                "location": "Kozhikode",
                "expense": "1250.50",
                "status": Visit.STATUS_INTERESTED,
                "note": "Discussed pricing.",
            },
        )

        visit = Visit.objects.get(customer=self.user_one_customer)
        self.assertRedirects(response, reverse("customers:visit_detail", args=[visit.id]))
        self.assertEqual(visit.location, "Kozhikode")
        self.assertEqual(visit.expense, Decimal("1250.50"))
        self.assertEqual(visit.status, Visit.STATUS_INTERESTED)

    def test_create_manual_client_visit(self):
        self.client.force_login(self.user_one)

        response = self.client.post(
            reverse("customers:visit_add"),
            {
                "visit_mode": "manual",
                "client_name": "Manual Prospect",
                "phone": "9999999999",
                "email": "manual@example.com",
                "company": "Manual Co",
                "client_location": "Kochi",
                "visit_date": "2026-10-02",
                "meeting_time": "10:00",
                "location": "Client office",
                "expense": "300.00",
                "status": Visit.STATUS_HEALTHY,
                "note": "Created manually.",
            },
        )

        manual_client = ManualClient.objects.get(client_name="Manual Prospect")
        visit = Visit.objects.get(manual_client=manual_client)
        self.assertRedirects(response, reverse("customers:visit_detail", args=[visit.id]))
        self.assertEqual(manual_client.owner, self.user_one)
        self.assertIsNone(visit.customer)
        self.assertEqual(visit.location, "Client office")

    def test_manual_visit_detail_shows_manual_client_data(self):
        self.client.force_login(self.user_one)
        manual_client = ManualClient.objects.create(
            owner=self.user_one,
            client_name="Walk In Client",
            phone="12345",
            company="Walk In Co",
            location="Kannur",
        )
        visit = Visit.objects.create(
            manual_client=manual_client,
            visit_date=date(2026, 10, 2),
            meeting_time=time(10, 0),
            location="Kannur",
            expense=Decimal("50.00"),
            status=Visit.STATUS_HEALTHY,
        )

        response = self.client.get(reverse("customers:visit_detail", args=[visit.id]))

        self.assertContains(response, "Manual Client Data")
        self.assertContains(response, "Walk In Client")
        self.assertContains(response, "Kannur")

    def test_user_ownership_blocks_other_users_manual_visit(self):
        manual_client = ManualClient.objects.create(
            owner=self.user_two,
            client_name="Other Manual Client",
        )
        visit = Visit.objects.create(
            manual_client=manual_client,
            visit_date=date(2026, 10, 2),
            meeting_time=time(10, 0),
            location="Delhi",
            expense=Decimal("50.00"),
            status=Visit.STATUS_HEALTHY,
        )
        self.client.force_login(self.user_one)

        detail_response = self.client.get(reverse("customers:visit_detail", args=[visit.id]))
        delete_response = self.client.post(reverse("customers:visit_delete", args=[visit.id]))

        self.assertEqual(detail_response.status_code, 404)
        self.assertEqual(delete_response.status_code, 404)
        self.assertTrue(Visit.objects.filter(id=visit.id).exists())

    def test_admin_can_access_manual_visits(self):
        manual_client = ManualClient.objects.create(
            owner=self.user_two,
            client_name="Admin Visible Manual Client",
        )
        Visit.objects.create(
            manual_client=manual_client,
            visit_date=date(2026, 10, 2),
            meeting_time=time(10, 0),
            location="Delhi",
            expense=Decimal("50.00"),
            status=Visit.STATUS_HEALTHY,
        )
        self.client.force_login(self.admin)

        response = self.client.get(reverse("customers:sales_list"))

        self.assertContains(response, "Admin Visible Manual Client")
        self.assertContains(response, "Manual Client")

    def test_edit_visit(self):
        self.client.force_login(self.user_one)
        visit = Visit.objects.create(
            customer=self.user_one_customer,
            visit_date=date(2026, 9, 30),
            meeting_time=time(9, 0),
            location="Initial",
            expense=Decimal("100.00"),
            status=Visit.STATUS_HEALTHY,
            note="Initial note",
        )

        response = self.client.post(
            reverse("customers:visit_edit", args=[visit.id]),
            {
                "dataset": self.user_one_dataset.id,
                "customer": self.user_one_customer.id,
                "customer_search": "",
                "visit_date": "2026-10-01",
                "meeting_time": "11:15",
                "location": "Updated location",
                "expense": "250.00",
                "status": Visit.STATUS_FOLLOW_UP,
                "note": "Follow up next week.",
            },
        )

        self.assertRedirects(response, reverse("customers:visit_detail", args=[visit.id]))
        visit.refresh_from_db()
        self.assertEqual(visit.location, "Updated location")
        self.assertEqual(visit.status, Visit.STATUS_FOLLOW_UP)
        self.assertEqual(visit.note, "Follow up next week.")

    def test_delete_visit(self):
        self.client.force_login(self.user_one)
        visit = Visit.objects.create(
            customer=self.user_one_customer,
            visit_date=date(2026, 9, 30),
            meeting_time=time(9, 0),
            location="Kozhikode",
            expense=Decimal("100.00"),
            status=Visit.STATUS_HEALTHY,
        )

        response = self.client.post(reverse("customers:visit_delete", args=[visit.id]))

        self.assertRedirects(response, reverse("customers:sales_list"))
        self.assertFalse(Visit.objects.filter(id=visit.id).exists())

    def test_user_ownership_blocks_other_users_visits(self):
        self.client.force_login(self.user_one)

        detail_response = self.client.get(reverse("customers:visit_detail", args=[self.user_two_visit.id]))
        edit_response = self.client.post(
            reverse("customers:visit_edit", args=[self.user_two_visit.id]),
            {
                "dataset": self.user_one_dataset.id,
                "customer": self.user_one_customer.id,
                "customer_search": "",
                "visit_date": "2026-09-30",
                "meeting_time": "12:00",
                "location": "Blocked",
                "expense": "1.00",
                "status": Visit.STATUS_CONVERTED,
            },
        )
        delete_response = self.client.post(reverse("customers:visit_delete", args=[self.user_two_visit.id]))

        self.assertEqual(detail_response.status_code, 404)
        self.assertEqual(edit_response.status_code, 404)
        self.assertEqual(delete_response.status_code, 404)
        self.assertTrue(Visit.objects.filter(id=self.user_two_visit.id).exists())

    def test_admin_can_access_all_visits(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse("customers:sales_list"))

        self.assertContains(response, "Asha Rao")
        self.assertContains(response, "User Two Customers")

    def test_dashboard_sales_statistics_are_scoped_to_user(self):
        Visit.objects.create(
            customer=self.user_one_customer,
            visit_date=date(2026, 9, 30),
            meeting_time=time(9, 0),
            location="Kozhikode",
            expense=Decimal("300.00"),
            status=Visit.STATUS_CONVERTED,
        )
        self.client.force_login(self.user_one)

        response = self.client.get(reverse("customers:dashboard"))

        self.assertContains(response, "Sales Visits")
        self.assertContains(response, "Total Visits")
        self.assertContains(response, "₹300.00")


class CRMValidationTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="validationuser",
            password="test-password-123",
        )
        self.client.force_login(self.user)

    def manual_visit_payload(self, **overrides):
        data = {
            "visit_mode": "manual",
            "client_name": "Valid Client",
            "phone": "+91 98765 43210",
            "email": "valid@example.com",
            "company": "Valid Company",
            "client_location": "Kochi",
            "visit_date": "2026-10-02",
            "meeting_time": "10:00",
            "location": "Client office",
            "expense": "300.00",
            "status": Visit.STATUS_HEALTHY,
            "note": "Normal visit note.",
        }
        data.update(overrides)
        return data

    def post_manual_visit(self, **overrides):
        return self.client.post(
            reverse("customers:visit_add"),
            self.manual_visit_payload(**overrides),
        )

    def test_valid_manual_client_and_visit_are_accepted(self):
        response = self.post_manual_visit()

        manual_client = ManualClient.objects.get(client_name="Valid Client")
        visit = Visit.objects.get(manual_client=manual_client)
        self.assertRedirects(response, reverse("customers:visit_detail", args=[visit.id]))
        self.assertEqual(manual_client.phone, "+91 98765 43210")
        self.assertEqual(manual_client.email, "valid@example.com")

    def test_invalid_phone_is_rejected(self):
        response = self.post_manual_visit(phone="phone-number")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Please enter a valid phone number.")
        self.assertEqual(Visit.objects.count(), 0)

    def test_valid_phone_with_country_code_is_accepted(self):
        response = self.post_manual_visit(phone="+91 98765-43210")

        manual_client = ManualClient.objects.get(client_name="Valid Client")
        visit = Visit.objects.get(manual_client=manual_client)
        self.assertRedirects(response, reverse("customers:visit_detail", args=[visit.id]))
        self.assertEqual(manual_client.phone, "+91 98765-43210")

    def test_invalid_email_is_rejected(self):
        response = self.post_manual_visit(email="name@")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Enter a valid email address.")
        self.assertEqual(Visit.objects.count(), 0)

    def test_valid_email_is_accepted(self):
        response = self.post_manual_visit(email="name@example.com")

        manual_client = ManualClient.objects.get(client_name="Valid Client")
        visit = Visit.objects.get(manual_client=manual_client)
        self.assertRedirects(response, reverse("customers:visit_detail", args=[visit.id]))
        self.assertEqual(manual_client.email, "name@example.com")

    def test_empty_required_fields_are_rejected(self):
        response = self.post_manual_visit(
            client_name="   ",
            visit_date="",
            meeting_time="",
            location="",
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Client Name is required.")
        self.assertContains(response, "This field is required.")
        self.assertEqual(Visit.objects.count(), 0)

    def test_invalid_expense_is_rejected(self):
        response = self.post_manual_visit(expense="12abc")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Enter a number.")
        self.assertEqual(Visit.objects.count(), 0)

    def test_negative_expense_is_rejected(self):
        response = self.post_manual_visit(expense="-100")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Expense cannot be negative.")
        self.assertEqual(Visit.objects.count(), 0)

    def test_invalid_date_is_rejected(self):
        response = self.post_manual_visit(visit_date="not-a-date")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Enter a valid date.")
        self.assertEqual(Visit.objects.count(), 0)

    def test_invalid_time_is_rejected(self):
        response = self.post_manual_visit(meeting_time="not-a-time")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Enter a valid time.")
        self.assertEqual(Visit.objects.count(), 0)

    def test_invalid_status_is_rejected(self):
        response = self.post_manual_visit(status="made_up_status")

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Select a valid choice.")
        self.assertEqual(Visit.objects.count(), 0)

    def test_numbers_only_and_symbol_only_manual_fields_are_rejected(self):
        response = self.post_manual_visit(
            client_name="123456",
            company="12345",
            client_location="@@@",
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Please enter a valid client name.")
        self.assertContains(response, "Please enter a valid company.")
        self.assertContains(response, "Please enter a valid location.")
        self.assertEqual(Visit.objects.count(), 0)

    def test_unreasonably_long_note_is_rejected(self):
        response = self.post_manual_visit(note="x" * 2001)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Note cannot be longer than 2000 characters.")
        self.assertEqual(Visit.objects.count(), 0)

    def test_empty_dataset_name_is_rejected(self):
        response = self.client.post(
            reverse("customers:import_excel"),
            {
                "dataset_name": "   ",
                "excel_file": SimpleUploadedFile(
                    "valid.xlsx",
                    build_workbook_bytes("Leads", ["Name"], [["Asha"]]),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Dataset name is required.")
        self.assertFalse(Dataset.objects.exists())

    def test_invalid_excel_upload_is_rejected(self):
        response = self.client.post(
            reverse("customers:import_excel"),
            {
                "dataset_name": "Valid Leads",
                "excel_file": SimpleUploadedFile(
                    "not-excel.xlsx",
                    b"this is not an excel workbook",
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "The uploaded file does not appear to be a valid .xlsx workbook.")
        self.assertFalse(Dataset.objects.exists())


class APIFoundationTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.user_one = User.objects.create_user(
            username="apiuserone",
            password="test-password-123",
        )
        self.user_two = User.objects.create_user(
            username="apiusertwo",
            password="test-password-123",
        )
        self.admin = User.objects.create_user(
            username="apiadmin",
            password="test-password-123",
            is_staff=True,
        )

        self.client.force_login(self.user_one)
        self.user_one_dataset = import_dataset(
            self.client,
            dataset_name="API User One Customers",
            filename="api-one.xlsx",
            headers=["Customer Name", "City"],
            rows=[["One Customer", "Kochi"]],
        )

        self.client.force_login(self.user_two)
        self.user_two_dataset = import_dataset(
            self.client,
            dataset_name="API User Two Customers",
            filename="api-two.xlsx",
            headers=["Customer Name", "City"],
            rows=[["Two Customer", "Delhi"]],
        )

    def test_session_endpoint_reports_authenticated_user(self):
        self.client.force_login(self.user_one)

        response = self.client.get(reverse("customers_api:session"))

        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["isAuthenticated"])
        self.assertEqual(response.json()["user"]["username"], "apiuserone")

    def test_react_login_api_uses_django_session_and_logout(self):
        self.client.logout()
        session = self.client.get(reverse("customers_api:session"))
        self.assertEqual(session.status_code, 200)
        response = self.client.post(
            reverse("customers_api:login"),
            data=json.dumps({"username": "apiuserone", "password": "test-password-123"}),
            content_type="application/json",
            HTTP_X_CSRFTOKEN=session.json()["csrfToken"],
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["user"]["username"], "apiuserone")
        self.assertTrue(self.client.get(reverse("customers_api:session")).json()["isAuthenticated"])
        logout_response = self.client.post(
            reverse("customers_api:logout"),
            data="{}",
            content_type="application/json",
        )
        self.assertEqual(logout_response.status_code, 204)
        self.assertFalse(self.client.get(reverse("customers_api:session")).json()["isAuthenticated"])

    def test_reminder_edit_is_owner_scoped_and_resets_delivery_when_schedule_changes(self):
        self.client.force_login(self.user_one)
        sent_at = timezone.now() - timedelta(minutes=5)
        reminder = Reminder.objects.create(
            owner=self.user_one,
            text="Original reminder",
            reminder_at=timezone.now() + timedelta(days=1),
            sent_at=sent_at,
        )
        new_time = timezone.now() + timedelta(days=2)
        response = self.client.patch(
            reverse("customers_api:reminder_detail", args=[reminder.id]),
            data=json.dumps({"text": "Updated reminder", "reminder_at": new_time.isoformat()}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        reminder.refresh_from_db()
        self.assertEqual(reminder.text, "Updated reminder")
        self.assertIsNone(reminder.sent_at)

        self.client.force_login(self.user_two)
        denied = self.client.patch(
            reverse("customers_api:reminder_detail", args=[reminder.id]),
            data=json.dumps({"text": "Other user's edit"}),
            content_type="application/json",
        )
        self.assertEqual(denied.status_code, 404)

    def test_reminder_template_can_open_and_save_an_edit(self):
        self.client.force_login(self.user_one)
        reminder = Reminder.objects.create(
            owner=self.user_one,
            text="Original reminder",
            reminder_at=timezone.now() + timedelta(days=1),
        )
        edit_page = self.client.get(reverse("customers:notifications"), {"edit": reminder.id})
        self.assertContains(edit_page, "Edit reminder")
        target = timezone.localtime(timezone.now() + timedelta(days=3))
        response = self.client.post(
            reverse("customers:notifications"),
            {
                "action": "update",
                "reminder_id": reminder.id,
                "text": "Edited reminder",
                "date": target.date().isoformat(),
                "time": target.strftime("%H:%M"),
            },
        )

        self.assertRedirects(response, reverse("customers:notifications"))
        reminder.refresh_from_db()
        self.assertEqual(reminder.text, "Edited reminder")

    def test_admin_can_assign_task_and_user_can_list_it(self):
        self.client.force_login(self.admin)
        due_at = timezone.now() + timedelta(days=1)

        response = self.client.post(
            reverse("customers_api:admin_task_list"),
            data=json.dumps({
                "assignee": self.user_one.id,
                "title": "Call converted leads",
                "note": "Check the hottest customers first.",
                "due_at": due_at.isoformat(),
            }),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.json()["title"], "Call converted leads")
        self.assertEqual(response.json()["assignee_username"], "apiuserone")

        self.client.force_login(self.user_one)
        list_response = self.client.get(reverse("customers_api:task_list"))

        self.assertEqual(list_response.status_code, 200)
        tasks = list_response.json()["results"]
        self.assertEqual(len(tasks), 1)
        self.assertEqual(tasks[0]["title"], "Call converted leads")
        self.assertEqual(tasks[0]["assigned_by_username"], "apiadmin")

    def test_user_can_only_update_own_task_status(self):
        own_task = AssignedTask.objects.create(
            assignee=self.user_one,
            assigned_by=self.admin,
            title="Own task",
        )
        other_task = AssignedTask.objects.create(
            assignee=self.user_two,
            assigned_by=self.admin,
            title="Other task",
        )
        self.client.force_login(self.user_one)

        response = self.client.patch(
            reverse("customers_api:task_detail", args=[own_task.id]),
            data=json.dumps({"status": "done"}),
            content_type="application/json",
        )
        denied = self.client.patch(
            reverse("customers_api:task_detail", args=[other_task.id]),
            data=json.dumps({"status": "done"}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(denied.status_code, 404)
        own_task.refresh_from_db()
        other_task.refresh_from_db()
        self.assertEqual(own_task.status, "done")
        self.assertEqual(other_task.status, "pending")

    def test_normal_user_cannot_access_admin_tasks_api(self):
        self.client.force_login(self.user_one)

        response = self.client.get(reverse("customers_api:admin_task_list"))

        self.assertEqual(response.status_code, 403)

    def test_react_login_api_rejects_invalid_and_inactive_users(self):
        self.client.logout()
        response = self.client.post(
            reverse("customers_api:login"),
            data=json.dumps({"username": "apiuserone", "password": "incorrect"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json()["detail"], "Invalid username or password.")
        self.user_one.is_active = False
        self.user_one.save(update_fields=["is_active"])
        response = self.client.post(
            reverse("customers_api:login"),
            data=json.dumps({"username": "apiuserone", "password": "test-password-123"}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 400)
        self.assertFalse(self.client.get(reverse("customers_api:session")).json()["isAuthenticated"])

    def test_login_api_enforces_csrf(self):
        csrf_client = self.client_class(enforce_csrf_checks=True)
        session = csrf_client.get(reverse("customers_api:session"))
        url = reverse("customers_api:login")
        payload = {"username": "apiuserone", "password": "test-password-123"}
        denied = csrf_client.post(url, data=payload)
        self.assertEqual(denied.status_code, 403)
        allowed = csrf_client.post(url, data=payload, HTTP_X_CSRFTOKEN=session.json()["csrfToken"])
        self.assertEqual(allowed.status_code, 200)

    def test_dataset_api_is_scoped_to_owner(self):
        self.client.force_login(self.user_one)

        response = self.client.get(reverse("customers_api:dataset_list"))

        self.assertEqual(response.status_code, 200)
        dataset_names = [dataset["name"] for dataset in response.json()["results"]]
        self.assertIn("API User One Customers", dataset_names)
        self.assertNotIn("API User Two Customers", dataset_names)

    def test_admin_dataset_api_can_see_all_datasets(self):
        self.client.force_login(self.admin)

        response = self.client.get(reverse("customers_api:dataset_list"))

        self.assertEqual(response.status_code, 200)
        dataset_names = [dataset["name"] for dataset in response.json()["results"]]
        self.assertIn("API User One Customers", dataset_names)
        self.assertIn("API User Two Customers", dataset_names)

    def test_dataset_api_includes_assigned_target_and_remaining_count(self):
        record = self.user_one_dataset.records.get()
        SalesTarget.objects.create(user=self.user_one, month=timezone.localdate().replace(day=1), target=10)
        for index in range(3):
            Visit.objects.create(
                customer=record,
                visit_date=timezone.localdate(),
                meeting_time=time(10 + index, 0),
                location="Kochi",
                expense=Decimal("0.00"),
                status=Visit.STATUS_CONVERTED,
            )
        self.client.force_login(self.user_one)

        response = self.client.get(reverse("customers_api:dataset_list"))

        self.assertEqual(response.status_code, 200)
        dataset = response.json()["results"][0]
        self.assertEqual(dataset["target"], 10)
        self.assertEqual(dataset["converted_count"], 3)
        self.assertEqual(dataset["remaining_target"], 7)

    def test_dataset_detail_api_blocks_other_users_dataset(self):
        self.client.force_login(self.user_one)

        response = self.client.get(
            reverse("customers_api:dataset_detail", args=[self.user_two_dataset.id])
        )

        self.assertEqual(response.status_code, 404)

    def test_dataset_management_api_enforces_ownership_and_validation(self):
        self.client.force_login(self.user_one)
        own_url = reverse("customers_api:dataset_manage", args=[self.user_one_dataset.id])
        invalid = self.client.patch(
            own_url,
            data=json.dumps({"name": "   "}),
            content_type="application/json",
        )
        self.assertEqual(invalid.status_code, 400)
        self.assertIn("name", invalid.json()["errors"])
        renamed = self.client.patch(
            own_url,
            data=json.dumps({"name": "Renamed Dataset"}),
            content_type="application/json",
        )
        self.assertEqual(renamed.status_code, 200)
        self.assertEqual(renamed.json()["name"], "Renamed Dataset")
        denied = self.client.delete(
            reverse("customers_api:dataset_manage", args=[self.user_two_dataset.id])
        )
        self.assertEqual(denied.status_code, 404)
        self.assertTrue(Dataset.objects.filter(pk=self.user_two_dataset.id).exists())

    def test_dataset_management_api_allows_admin_delete(self):
        self.client.force_login(self.admin)
        response = self.client.delete(
            reverse("customers_api:dataset_manage", args=[self.user_two_dataset.id])
        )
        self.assertEqual(response.status_code, 204)
        self.assertFalse(Dataset.objects.filter(pk=self.user_two_dataset.id).exists())

    def test_customer_by_id_api_is_owner_scoped(self):
        self.client.force_login(self.user_one)
        own_record = self.user_one_dataset.records.get()
        other_record = self.user_two_dataset.records.get()
        self.assertEqual(
            self.client.get(reverse("customers_api:customer_by_id", args=[own_record.id])).status_code,
            200,
        )
        self.assertEqual(
            self.client.get(reverse("customers_api:customer_by_id", args=[other_record.id])).status_code,
            404,
        )

    def test_dashboard_api_is_scoped_to_owner(self):
        self.client.force_login(self.user_one)

        response = self.client.get(reverse("customers_api:dashboard"))

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["totalDatasets"], 1)
        self.assertEqual(data["totalCustomers"], 1)
        recent_names = [dataset["name"] for dataset in data["recentDatasets"]]
        self.assertIn("API User One Customers", recent_names)
        self.assertNotIn("API User Two Customers", recent_names)

    def test_dataset_detail_api_returns_dynamic_columns_and_records(self):
        self.client.force_login(self.user_one)

        response = self.client.get(
            reverse("customers_api:dataset_detail", args=[self.user_one_dataset.id])
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["dataset"]["name"], "API User One Customers")
        self.assertEqual(data["recordsBySheet"][0]["sheetName"], "Leads")
        column_names = [
            column["originalName"]
            for column in data["recordsBySheet"][0]["columns"]
        ]
        self.assertEqual(column_names, ["Customer Name", "City"])
        self.assertEqual(data["recordsBySheet"][0]["records"][0]["values"], ["One Customer", "Kochi"])

    def test_dataset_detail_api_search_is_dataset_specific(self):
        self.client.force_login(self.user_one)
        record = self.user_one_dataset.records.get()
        record.original_data["city"] = "Delhi"
        record.save()

        response = self.client.get(
            reverse("customers_api:dataset_detail", args=[self.user_one_dataset.id]),
            {"q": "Delhi"},
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()
        values = [
            value
            for sheet in data["recordsBySheet"]
            for item in sheet["records"]
            for value in item["values"]
        ]
        self.assertIn("Delhi", values)
        self.assertNotIn("Two Customer", values)

    def test_dataset_detail_api_response_filter(self):
        self.client.force_login(self.user_one)
        record = self.user_one_dataset.records.get()
        record.response = CustomerRecord.RESPONSE_INTERESTED
        record.save()

        interested_response = self.client.get(
            reverse("customers_api:dataset_detail", args=[self.user_one_dataset.id]),
            {"response": CustomerRecord.RESPONSE_INTERESTED},
        )
        callback_response = self.client.get(
            reverse("customers_api:dataset_detail", args=[self.user_one_dataset.id]),
            {"response": CustomerRecord.RESPONSE_CALLBACK},
        )

        self.assertEqual(interested_response.status_code, 200)
        self.assertEqual(interested_response.json()["pagination"]["count"], 1)
        self.assertEqual(callback_response.json()["pagination"]["count"], 0)

    def test_dataset_detail_api_paginates_records(self):
        self.client.force_login(self.user_one)
        large_dataset = import_dataset(
            self.client,
            dataset_name="API Large Customers",
            filename="api-large.xlsx",
            headers=["Customer Name"],
            rows=[[f"Customer {index}"] for index in range(1, 26)],
        )

        first_page = self.client.get(
            reverse("customers_api:dataset_detail", args=[large_dataset.id])
        )
        second_page = self.client.get(
            reverse("customers_api:dataset_detail", args=[large_dataset.id]),
            {"page": 2},
        )

        self.assertEqual(first_page.json()["pagination"]["count"], 25)
        self.assertEqual(first_page.json()["pagination"]["numPages"], 2)
        self.assertEqual(len(first_page.json()["recordsBySheet"][0]["records"]), 20)
        self.assertEqual(len(second_page.json()["recordsBySheet"][0]["records"]), 5)

    def test_admin_dataset_detail_api_can_access_all_datasets(self):
        self.client.force_login(self.admin)

        response = self.client.get(
            reverse("customers_api:dataset_detail", args=[self.user_two_dataset.id])
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["dataset"]["name"], "API User Two Customers")

    def test_customer_detail_api_blocks_direct_id_access_to_other_user_customer(self):
        self.client.force_login(self.user_one)
        other_record = self.user_two_dataset.records.get()

        response = self.client.get(
            reverse(
                "customers_api:customer_detail",
                args=[self.user_two_dataset.id, other_record.id],
            )
        )

        self.assertEqual(response.status_code, 404)

    def test_customer_detail_api_returns_field_rows_and_crm_data(self):
        self.client.force_login(self.user_one)
        record = self.user_one_dataset.records.get()

        response = self.client.get(
            reverse("customers_api:customer_detail", args=[self.user_one_dataset.id, record.id])
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["dataset"]["name"], "API User One Customers")
        self.assertEqual(data["record"]["id"], record.id)
        labels = [field["label"] for field in data["fieldRows"]]
        self.assertEqual(labels, ["Customer Name", "City"])
        self.assertIn({"value": CustomerRecord.RESPONSE_CALLBACK, "label": "Callback"}, data["responseChoices"])

    def test_customer_update_api_updates_only_crm_fields(self):
        self.client.force_login(self.user_one)
        record = self.user_one_dataset.records.get()
        original_data = record.original_data.copy()

        response = self.client.patch(
            reverse("customers_api:customer_detail", args=[self.user_one_dataset.id, record.id]),
            data=json.dumps(
                {
                    "response": CustomerRecord.RESPONSE_CALLBACK,
                    "note": "Call again tomorrow.",
                }
            ),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 200)
        record.refresh_from_db()
        self.assertEqual(record.response, CustomerRecord.RESPONSE_CALLBACK)
        self.assertEqual(record.note, "Call again tomorrow.")
        self.assertEqual(record.original_data, original_data)

    def test_customer_update_api_rejects_invalid_response(self):
        self.client.force_login(self.user_one)
        record = self.user_one_dataset.records.get()

        response = self.client.patch(
            reverse("customers_api:customer_detail", args=[self.user_one_dataset.id, record.id]),
            data=json.dumps({"response": "invalid", "note": ""}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("response", response.json()["errors"])

    def test_customer_delete_api_deletes_only_accessible_customer(self):
        self.client.force_login(self.user_one)
        record = self.user_one_dataset.records.get()

        response = self.client.delete(
            reverse("customers_api:customer_detail", args=[self.user_one_dataset.id, record.id])
        )

        self.assertEqual(response.status_code, 204)
        self.assertFalse(CustomerRecord.objects.filter(id=record.id).exists())
        self.assertTrue(CustomerRecord.objects.filter(dataset=self.user_two_dataset).exists())

    def test_customer_delete_api_blocks_other_user_customer(self):
        self.client.force_login(self.user_one)
        other_record = self.user_two_dataset.records.get()

        response = self.client.delete(
            reverse("customers_api:customer_detail", args=[self.user_two_dataset.id, other_record.id])
        )

        self.assertEqual(response.status_code, 404)
        self.assertTrue(CustomerRecord.objects.filter(id=other_record.id).exists())

    def test_excel_inspect_api_returns_workbook_structure(self):
        self.client.force_login(self.user_one)

        response = self.client.post(
            reverse("customers_api:import_inspect"),
            {
                "dataset_name": "API Import Leads",
                "excel_file": SimpleUploadedFile(
                    "api-import.xlsx",
                    build_workbook_bytes("Leads", ["Customer", "Phone"], [["Asha", "9876543210"]]),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )

        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["datasetName"], "API Import Leads")
        self.assertEqual(data["inspection"]["filename"], "api-import.xlsx")
        self.assertEqual(data["inspection"]["sheets"][0]["column_names"], ["Customer", "Phone"])
        self.assertTrue(data["pendingImportToken"])

    def test_excel_import_confirm_api_creates_owned_dataset(self):
        self.client.force_login(self.user_one)
        inspect_response = self.client.post(
            reverse("customers_api:import_inspect"),
            {
                "dataset_name": "API Confirm Leads",
                "excel_file": SimpleUploadedFile(
                    "api-confirm.xlsx",
                    build_workbook_bytes("Leads", ["Customer"], [["Asha"]]),
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )
        token = inspect_response.json()["pendingImportToken"]

        response = self.client.post(
            reverse("customers_api:import_confirm"),
            data=json.dumps({"pendingImportToken": token}),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 201)
        dataset = Dataset.objects.get(name="API Confirm Leads")
        self.assertEqual(dataset.owner, self.user_one)
        self.assertEqual(dataset.records.count(), 1)
        self.assertEqual(response.json()["dataset"]["id"], dataset.id)

    def test_excel_inspect_api_rejects_invalid_upload(self):
        self.client.force_login(self.user_one)

        response = self.client.post(
            reverse("customers_api:import_inspect"),
            {
                "dataset_name": "Invalid Upload",
                "excel_file": SimpleUploadedFile(
                    "invalid.xlsx",
                    b"not an excel file",
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            },
        )

        self.assertEqual(response.status_code, 400)
        self.assertIn("excel_file", response.json()["errors"])

    def test_dataset_export_api_exports_filtered_results(self):
        self.client.force_login(self.user_one)
        dataset = import_dataset(
            self.client,
            dataset_name="API Export Leads",
            filename="api-export.xlsx",
            headers=["Customer Name", "City"],
            rows=[
                ["Rahul Nair", "Kochi"],
                ["Rahul Menon", "Delhi"],
                ["Meera Nair", "Kochi"],
            ],
        )
        interested = dataset.records.get(original_data__customer_name="Rahul Nair")
        interested.response = CustomerRecord.RESPONSE_INTERESTED
        interested.note = "Export this row."
        interested.save()

        response = self.client.get(
            reverse("customers_api:dataset_export", args=[dataset.id]),
            {"q": "Rahul", "response": CustomerRecord.RESPONSE_INTERESTED},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        rows = read_export_rows(response.content)
        self.assertEqual(rows[0], ["Customer Name", "City", "Response", "Note"])
        self.assertEqual(rows[1], ["Rahul Nair", "Kochi", "Interested", "Export this row."])
        self.assertEqual(len(rows), 2)

    def test_sales_api_lists_owned_visits_and_blocks_other_owner(self):
        own_record = self.user_one_dataset.records.get()
        other_record = self.user_two_dataset.records.get()
        own_visit = Visit.objects.create(
            customer=own_record, visit_date=date(2026, 10, 1), meeting_time=time(10, 0),
            location="Kochi", expense=Decimal("25.50"), status=Visit.STATUS_HEALTHY,
        )
        other_visit = Visit.objects.create(
            customer=other_record, visit_date=date(2026, 10, 1), meeting_time=time(11, 0),
            location="Delhi", expense=Decimal("40.00"), status=Visit.STATUS_INTERESTED,
        )
        self.client.force_login(self.user_one)
        response = self.client.get(reverse("customers_api:visit_list"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual([item["id"] for item in response.json()["results"]], [own_visit.id])
        self.assertEqual(
            self.client.get(reverse("customers_api:visit_detail", args=[other_visit.id])).status_code,
            404,
        )

    def test_sales_api_creates_manual_visit_for_logged_in_user(self):
        self.client.force_login(self.user_one)
        response = self.client.post(
            reverse("customers_api:visit_list"),
            data=json.dumps({
                "visit_mode": "manual", "client_name": "Manual Customer",
                "phone": "+91 98765 43210", "email": "manual@example.com",
                "company": "Example Co", "client_location": "Kochi",
                "visit_date": "2026-10-01", "meeting_time": "12:30",
                "location": "Office", "expense": "100.50", "status": Visit.STATUS_CONVERTED,
                "note": "Follow up next week",
            }), content_type="application/json",
        )
        self.assertEqual(response.status_code, 201)
        visit = Visit.objects.get(pk=response.json()["id"])
        self.assertEqual(visit.manual_client.owner, self.user_one)
        self.assertIsNone(visit.customer_id)

    def test_sales_api_client_options_are_dataset_scoped(self):
        self.client.force_login(self.user_one)
        own = self.client.get(reverse("customers_api:visit_clients"), {"dataset": self.user_one_dataset.id})
        self.assertEqual(own.status_code, 200)
        self.assertEqual(len(own.json()["results"]), 1)
        denied = self.client.get(reverse("customers_api:visit_clients"), {"dataset": self.user_two_dataset.id})
        self.assertEqual(denied.status_code, 404)

    def test_admin_sales_api_can_see_all_visits_and_admin_panel_is_protected(self):
        Visit.objects.create(
            customer=self.user_two_dataset.records.get(), visit_date=date(2026, 10, 1),
            meeting_time=time(10, 0), location="Delhi", expense=Decimal("10.00"),
            status=Visit.STATUS_HEALTHY,
        )
        self.client.force_login(self.admin)
        self.assertEqual(len(self.client.get(reverse("customers_api:visit_list")).json()["results"]), 1)
        self.assertEqual(self.client.get(reverse("customers_api:admin_summary")).status_code, 200)
        self.client.force_login(self.user_one)
        self.assertEqual(self.client.get(reverse("customers_api:admin_summary")).status_code, 403)

    def test_admin_target_api_assigns_target_and_reports_progress_warning(self):
        today = timezone.localdate()
        month = today.replace(day=1)
        previous_month = (month - timedelta(days=1)).replace(day=1)
        older_month = (previous_month - timedelta(days=1)).replace(day=1)
        record = self.user_one_dataset.records.get()

        SalesTarget.objects.create(user=self.user_one, month=previous_month, target=10)
        SalesTarget.objects.create(user=self.user_one, month=older_month, target=10)
        for index in range(3):
            Visit.objects.create(
                customer=record,
                visit_date=today,
                meeting_time=time(10 + index, 0),
                location="Kochi",
                expense=Decimal("0.00"),
                status=Visit.STATUS_CONVERTED,
            )

        self.client.force_login(self.admin)
        create_response = self.client.post(
            reverse("customers_api:admin-targets"),
            data=json.dumps({
                "userId": self.user_one.id,
                "month": month.isoformat()[:7],
                "target": 10,
            }),
            content_type="application/json",
        )
        list_response = self.client.get(reverse("customers_api:admin-targets"))

        self.assertEqual(create_response.status_code, 200)
        payload = next(
            item for item in list_response.json()
            if item["month"] == month.isoformat()[:7]
        )
        self.assertEqual(payload["userId"], self.user_one.id)
        self.assertEqual(payload["month"], month.isoformat()[:7])
        self.assertEqual(payload["monthLabel"], month.strftime("%B %Y"))
        self.assertEqual(payload["target"], 10)
        self.assertEqual(payload["converted"], 3)
        self.assertEqual(payload["percentage"], 30)
        self.assertEqual(payload["statusColor"], "danger")
        self.assertEqual(payload["lowTargetCount"], 3)
        self.assertTrue(payload["needsWarning"])
        self.assertIn("failed to achieve at least 30%", payload["warningMessage"])

    def test_admin_target_api_assigns_target_to_selected_user_month(self):
        self.client.force_login(self.admin)
        month = timezone.localdate().replace(day=1)

        create_response = self.client.post(
            reverse("customers_api:admin-targets"),
            data=json.dumps({
                "userId": self.user_two.id,
                "month": month.isoformat()[:7],
                "target": 15,
            }),
            content_type="application/json",
        )
        list_response = self.client.get(reverse("customers_api:admin-targets"))

        self.assertEqual(create_response.status_code, 200)
        self.assertEqual(create_response.json()["month"], month.isoformat()[:7])
        self.assertEqual(create_response.json()["monthLabel"], month.strftime("%B %Y"))
        self.assertEqual(len(list_response.json()), 1)
        self.assertEqual(list_response.json()[0]["userId"], self.user_two.id)

    def test_admin_target_api_rejects_invalid_month(self):
        self.client.force_login(self.admin)

        response = self.client.post(
            reverse("customers_api:admin-targets"),
            data=json.dumps({
                "userId": self.user_one.id,
                "month": "not-a-month",
                "target": 15,
            }),
            content_type="application/json",
        )

        self.assertEqual(response.status_code, 400)

    def test_user_target_api_lists_only_user_monthly_targets(self):
        current_month = timezone.localdate().replace(day=1)
        previous_month = (current_month - timedelta(days=1)).replace(day=1)
        next_month = (
            current_month.replace(year=current_month.year + 1, month=1, day=1)
            if current_month.month == 12
            else current_month.replace(month=current_month.month + 1, day=1)
        )
        SalesTarget.objects.create(user=self.user_one, month=previous_month, target=5)
        SalesTarget.objects.create(user=self.user_one, month=current_month, target=10)
        SalesTarget.objects.create(user=self.user_one, month=next_month, target=15)
        SalesTarget.objects.create(user=self.user_two, month=current_month, target=99)
        self.client.force_login(self.user_one)

        response = self.client.get(reverse("customers_api:target_list"))

        self.assertEqual(response.status_code, 200)
        targets = response.json()["results"]
        self.assertEqual(len(targets), 3)
        self.assertEqual(
            {target["month"] for target in targets},
            {
                previous_month.isoformat()[:7],
                current_month.isoformat()[:7],
                next_month.isoformat()[:7],
            },
        )
        self.assertNotIn(99, {target["target"] for target in targets})

    def test_admin_user_datasets_api_is_not_captured_by_action_route(self):
        self.client.force_login(self.admin)

        response = self.client.get(
            reverse("customers_api:admin_user_datasets", args=[self.user_one.id])
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["user"]["id"], self.user_one.id)
        self.assertEqual(len(response.json()["results"]), 1)


def import_dataset(client, dataset_name, filename, headers, rows):
    inspect_response = client.post(
        reverse("customers:import_excel"),
        {
            "dataset_name": dataset_name,
            "excel_file": SimpleUploadedFile(
                filename,
                build_workbook_bytes("Leads", headers, rows),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
    )
    assert inspect_response.status_code == 200
    assert Dataset.objects.filter(name=dataset_name).count() == 0
    token = inspect_response.context["pending_import_token"]

    response = client.post(
        reverse("customers:import_excel"),
        {
            "action": "confirm_import",
            "pending_import_token": token,
        },
    )
    dataset = Dataset.objects.get(name=dataset_name)
    assert response.status_code == 302
    return dataset


def build_workbook_bytes(sheet_name, headers, rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(headers)
    for row in rows:
        worksheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.read()


def read_export_rows(content):
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    workbook.close()
    return rows

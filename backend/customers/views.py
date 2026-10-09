import io
import os
import re
import uuid
import zipfile
from pathlib import Path

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.conf import settings
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Count, Prefetch, Q, Sum
from django.core.paginator import Paginator
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .forms import (
    AdminPasswordResetForm,
    AdminUserCreateForm,
    AdminUserEditForm,
    CustomerCRMForm,
    DatasetForm,
    ExcelUploadForm,
    ManualVisitForm,
    ReminderForm,
    save_visit_images,
    VisitForm,
)
from .models import AssignedTask, CustomerRecord, Dataset, DatasetColumn, ManualClient, Reminder, Visit


EMAIL_PATTERN = re.compile(r"([A-Z0-9._%+-]+)@([A-Z0-9.-]+\.[A-Z]{2,})", re.IGNORECASE)
DIGIT_PATTERN = re.compile(r"\d")


def crm_admin_required(view_func):
    @login_required
    def wrapped(request, *args, **kwargs):
        if not user_can_access_all_datasets(request.user):
            return HttpResponseForbidden("You do not have permission to access the admin panel.")
        return view_func(request, *args, **kwargs)

    return wrapped


@login_required
def dashboard(request):
    datasets = get_accessible_datasets(request.user)
    customer_records = CustomerRecord.objects.filter(dataset__in=datasets)
    visits = get_accessible_visits(request.user)
    today = timezone.localdate()
    show_all_datasets = request.GET.get("datasets") == "all"
    dataset_query = datasets.annotate(customer_count=Count("records"))
    if show_all_datasets:
        dataset_page = Paginator(dataset_query, 5).get_page(request.GET.get("page", 1))
        recent_datasets = dataset_page.object_list
    else:
        dataset_page = None
        recent_datasets = dataset_query[:3]
    response_counts = {
        row["response"]: row["count"]
        for row in customer_records.values("response").annotate(count=Count("id"))
    }
    response_stats = [
        {
            "value": value,
            "label": label,
            "count": response_counts.get(value, 0),
        }
        for value, label in CustomerRecord.RESPONSE_CHOICES
    ]
    return render(
        request,
        "customers/dashboard.html",
        {
            "total_datasets": datasets.count(),
            "total_customers": customer_records.count(),
            "recent_datasets": recent_datasets,
            "show_all_datasets": show_all_datasets,
            "dataset_page": dataset_page,
            "response_stats": response_stats,
            "total_visits": visits.count(),
            "todays_visits": visits.filter(visit_date=today).count(),
            "month_visits": visits.filter(visit_date__year=today.year, visit_date__month=today.month).count(),
            "total_visit_expenses": visits.aggregate(total=Sum("expense"))["total"] or 0,
        },
    )


@login_required
def sales_list(request):
    status_filter = request.GET.get("status", "").strip()
    date_from_value = request.GET.get("date_from", "").strip()
    date_to_value = request.GET.get("date_to", "").strip()
    try:
        date_from = parse_date(date_from_value) if date_from_value else None
    except ValueError:
        date_from = None
    try:
        date_to = parse_date(date_to_value) if date_to_value else None
    except ValueError:
        date_to = None
    visits = get_accessible_visits(request.user).select_related(
        "customer",
        "customer__dataset",
        "manual_client",
        "manual_client__owner",
    )
    if status_filter in dict(Visit.STATUS_CHOICES):
        visits = visits.filter(status=status_filter)
    if date_from:
        visits = visits.filter(visit_date__gte=date_from)
    if date_to:
        visits = visits.filter(visit_date__lte=date_to)
    total_visits = get_accessible_visits(request.user).count()
    visit_page = Paginator(visits, 5).get_page(request.GET.get("page", 1))
    return render(
        request,
        "customers/sales_list.html",
        {
            "visits": visit_page,
            "status_choices": Visit.STATUS_CHOICES,
            "status_filter": status_filter,
            "date_from": date_from_value,
            "date_to": date_to_value,
            "total_visits": total_visits,
            "filtered_visit_count": visit_page.paginator.count,
            "filters_active": bool(status_filter or date_from_value or date_to_value),
        },
    )


@login_required
def reminders(request):
    editing_reminder = None
    edit_id = request.GET.get("edit", "") if request.method == "GET" else (
        request.POST.get("reminder_id", "") if request.POST.get("action") == "update" else ""
    )
    if edit_id.isdigit():
        editing_reminder = get_object_or_404(Reminder, pk=edit_id, owner=request.user)
    initial = None
    if editing_reminder:
        reminder_time = timezone.localtime(editing_reminder.reminder_at)
        initial = {"text": editing_reminder.text, "date": reminder_time.date(), "time": reminder_time.time()}
    form = ReminderForm(request.POST or None, initial=initial)
    if request.method == "POST":
        action = request.POST.get("action", "create")
        if action == "cancel_edit":
            return redirect("customers:reminders")
        if action == "delete":
            try:
                reminder_id = int(request.POST.get("reminder_id", ""))
            except (TypeError, ValueError):
                reminder_id = None
            if reminder_id is not None:
                Reminder.objects.filter(pk=reminder_id, owner=request.user).delete()
            messages.success(request, "Reminder removed.")
            return redirect("customers:reminders")
        if form.is_valid():
            if action == "update":
                try:
                    reminder_id = int(request.POST.get("reminder_id", ""))
                except (TypeError, ValueError):
                    reminder_id = None
                reminder = get_object_or_404(Reminder, pk=reminder_id, owner=request.user)
                form.save(request.user, instance=reminder)
                messages.success(request, "Reminder updated successfully.")
            else:
                form.save(request.user)
                messages.success(request, "Reminder set successfully.")
            return redirect("customers:reminders")

    now = timezone.localtime()
    reminders = list(Reminder.objects.filter(owner=request.user))
    due_reminders = [reminder for reminder in reminders if reminder.is_due]
    upcoming_reminders = [reminder for reminder in reminders if not reminder.is_due]
    return render(
        request,
        "customers/reminders.html",
        {
            "form": form,
            "editing_reminder": editing_reminder,
            "reminders": reminders,
            "due_reminders": due_reminders,
            "upcoming_reminders": upcoming_reminders,
            "today": now.date().isoformat(),
            "current_time": now.strftime("%H:%M"),
        },
    )


@login_required
def notifications(request):
    if request.method == "GET" and ("reminder" in request.GET or "edit" in request.GET):
        query = request.META.get("QUERY_STRING", "")
        target = "customers:reminders"
        if query:
            return redirect(f"{redirect(target).url}?{query}")
        return redirect(target)

    if request.method == "POST" and request.POST.get("action") == "complete_task":
        task_id = request.POST.get("task_id", "")
        if task_id.isdigit():
            AssignedTask.objects.filter(pk=task_id, assignee=request.user).update(status=AssignedTask.STATUS_DONE)
            messages.success(request, "Task marked as completed.")
        return redirect("customers:notifications")

    tasks = list(AssignedTask.objects.filter(assignee=request.user).select_related("assigned_by"))
    pending_tasks = [task for task in tasks if task.status == AssignedTask.STATUS_PENDING]
    completed_tasks = [task for task in tasks if task.status == AssignedTask.STATUS_DONE]
    return render(
        request,
        "customers/notifications.html",
        {
            "tasks": tasks,
            "pending_tasks": pending_tasks,
            "completed_tasks": completed_tasks,
        },
    )


@login_required
def sales_clients(request):
    dataset = get_selected_dataset(request, get_accessible_datasets(request.user))
    if not dataset:
        return JsonResponse({"clients": []})

    customer_search = request.GET.get("q", "").strip()
    customers = get_visit_customer_options(dataset, customer_search)
    return JsonResponse(
        {
            "clients": [
                {
                    "id": customer.id,
                    "label": customer.display_name,
                    "sheet_name": customer.sheet_name,
                    "row_number": customer.row_number,
                }
                for customer in customers
            ]
        }
    )


@login_required
def visit_add(request):
    visit_mode = request.POST.get("visit_mode") if request.method == "POST" else request.GET.get("mode", "existing")
    if visit_mode not in ("existing", "manual"):
        visit_mode = "existing"
    datasets = get_accessible_datasets(request.user)
    selected_dataset = get_selected_dataset(request, datasets)
    customer_search = request.GET.get("customer_q", "")
    if request.method == "POST" and visit_mode == "existing":
        customer_search = request.POST.get("customer_search", "")
        selected_dataset = get_selected_dataset(request, datasets, source=request.POST)

    if visit_mode == "manual":
        form = ManualVisitForm(request.POST or None, request.FILES or None, owner=request.user)
    else:
        form = VisitForm(
            request.POST or None,
            request.FILES or None,
            datasets=datasets,
            selected_dataset=selected_dataset,
            customer_search=customer_search,
        )

    if request.method == "POST" and form.is_valid():
        visit = form.save()
        save_visit_images(visit, form.cleaned_data["images"])
        messages.success(request, "Client visit created successfully.")
        return redirect("customers:visit_detail", visit_id=visit.id)

    return render(
        request,
        "customers/visit_form.html",
        {
            "form": form,
            "page_title": "Add Client Visit",
            "page_subtitle": "Record a client visit for a customer in one of your datasets.",
            "submit_label": "Save Visit",
            "selected_dataset": selected_dataset,
            "customer_search": customer_search,
            "is_edit": False,
            "visit_mode": visit_mode,
        },
    )


@login_required
def visit_detail(request, visit_id):
    visit = get_accessible_visit_or_404(request.user, visit_id)
    if visit.customer_id:
        field_rows = build_field_rows(visit.customer, get_record_columns(visit.customer.dataset, visit.customer))
        client_heading = "Imported Customer Data"
    else:
        field_rows = build_manual_client_rows(visit.manual_client)
        client_heading = "Manual Client Data"
    return render(
        request,
        "customers/visit_detail.html",
        {
            "visit": visit,
            "visit_images": visit.images.all(),
            "field_rows": field_rows,
            "client_heading": client_heading,
        },
    )


@login_required
def visit_edit(request, visit_id):
    visit = get_accessible_visit_or_404(request.user, visit_id)
    datasets = get_accessible_datasets(request.user)
    if visit.manual_client_id:
        visit_mode = "manual"
        selected_dataset = None
        customer_search = ""
        form = ManualVisitForm(
            request.POST or None,
            request.FILES or None,
            instance=visit,
            owner=visit.manual_client.owner,
        )
    else:
        visit_mode = "existing"
        selected_dataset = get_selected_dataset(request, datasets, source=request.POST) if request.method == "POST" else visit.customer.dataset
        customer_search = request.POST.get("customer_search", "") if request.method == "POST" else request.GET.get("customer_q", "")
        form = VisitForm(
            request.POST or None,
            request.FILES or None,
            instance=visit,
            datasets=datasets,
            selected_dataset=selected_dataset,
            customer_search=customer_search,
        )

    if request.method == "POST" and form.is_valid():
        visit = form.save()
        save_visit_images(visit, form.cleaned_data["images"])
        messages.success(request, "Client visit updated successfully.")
        return redirect("customers:visit_detail", visit_id=visit.id)

    return render(
        request,
        "customers/visit_form.html",
        {
            "form": form,
            "visit": visit,
            "page_title": "Edit Client Visit",
            "page_subtitle": "Update visit details without changing imported customer data.",
            "submit_label": "Save Changes",
            "selected_dataset": selected_dataset,
            "customer_search": customer_search,
            "is_edit": True,
            "visit_mode": visit_mode,
        },
    )


@login_required
def visit_delete(request, visit_id):
    visit = get_accessible_visit_or_404(request.user, visit_id)

    if request.method == "POST":
        visit.delete()
        messages.success(request, "Client visit deleted successfully.")
        return redirect("customers:sales_list")

    return render(
        request,
        "customers/visit_confirm_delete.html",
        {
            "visit": visit,
        },
    )


@crm_admin_required
def admin_panel(request):
    User = get_user_model()
    users = Paginator(User.objects.annotate(
        dataset_count=Count("datasets", distinct=True),
        customer_count=Count("datasets__records", distinct=True),
    ).order_by("username"), 5).get_page(request.GET.get("page", 1))
    recent_datasets = Dataset.objects.select_related("owner").annotate(
        customer_count=Count("records")
    )[:5]
    recent_users = User.objects.order_by("-date_joined")[:5]

    return render(
        request,
        "customers/admin_panel.html",
        {
            "users": users,
            "total_users": User.objects.count(),
            "active_users": User.objects.filter(is_active=True).count(),
            "total_datasets": Dataset.objects.count(),
            "total_customers": CustomerRecord.objects.count(),
            "recent_datasets": recent_datasets,
            "recent_users": recent_users,
        },
    )


@crm_admin_required
def admin_user_add(request):
    form = AdminUserCreateForm(request.POST or None)

    if request.method == "POST" and form.is_valid():
        user = form.save()
        messages.success(request, f"User '{user.username}' created successfully.")
        return redirect("customers:admin_user_detail", user_id=user.id)

    return render(
        request,
        "customers/admin_user_form.html",
        {
            "form": form,
            "page_title": "Add User",
            "page_subtitle": "Create a normal CRM user account.",
            "submit_label": "Create User",
        },
    )


@crm_admin_required
def admin_user_detail(request, user_id):
    user = get_object_or_404(get_admin_user_queryset(), pk=user_id)

    return render(
        request,
        "customers/admin_user_detail.html",
        {
            "managed_user": user,
        },
    )


@crm_admin_required
def admin_user_edit(request, user_id):
    user = get_object_or_404(get_admin_user_queryset(), pk=user_id)
    form = AdminUserEditForm(request.POST or None, instance=user)

    if request.method == "POST" and form.is_valid():
        if user == request.user and not form.cleaned_data["is_active"]:
            form.add_error("is_active", "You cannot deactivate your own administrator account.")
        else:
            form.save()
            messages.success(request, f"User '{user.username}' updated successfully.")
            return redirect("customers:admin_user_detail", user_id=user.id)

    return render(
        request,
        "customers/admin_user_form.html",
        {
            "form": form,
            "managed_user": user,
            "page_title": "Edit User",
            "page_subtitle": "Update account details and active status.",
            "submit_label": "Save Changes",
        },
    )


@crm_admin_required
def admin_user_password(request, user_id):
    user = get_object_or_404(get_admin_user_queryset(), pk=user_id)

    if user_can_access_all_datasets(user):
        return HttpResponseForbidden("Password reset is available for normal CRM users only.")

    form = AdminPasswordResetForm(request.POST or None, user=user)

    if request.method == "POST" and form.is_valid():
        user.set_password(form.cleaned_data["password1"])
        user.save(update_fields=["password"])
        messages.success(request, f"Password reset for '{user.username}'.")
        return redirect("customers:admin_user_detail", user_id=user.id)

    return render(
        request,
        "customers/admin_user_password.html",
        {
            "form": form,
            "managed_user": user,
        },
    )


@crm_admin_required
def admin_user_toggle_active(request, user_id):
    user = get_object_or_404(get_admin_user_queryset(), pk=user_id)

    if request.method != "POST":
        return HttpResponseForbidden("Active status changes require a POST request.")

    if user == request.user:
        messages.error(request, "You cannot change your own active status here.")
        return redirect("customers:admin_user_detail", user_id=user.id)

    user.is_active = not user.is_active
    user.save(update_fields=["is_active"])
    status = "activated" if user.is_active else "deactivated"
    messages.success(request, f"User '{user.username}' {status}.")
    return redirect("customers:admin_user_detail", user_id=user.id)


@crm_admin_required
def admin_user_datasets(request, user_id):
    user = get_object_or_404(get_admin_user_queryset(), pk=user_id)
    datasets = user.datasets.annotate(customer_count=Count("records")).order_by("-uploaded_at")

    return render(
        request,
        "customers/admin_user_datasets.html",
        {
            "managed_user": user,
            "datasets": datasets,
        },
    )


@login_required
def dataset_list(request):
    datasets = Paginator(
        get_accessible_datasets(request.user).annotate(customer_count=Count("records")),
        5,
    ).get_page(request.GET.get("page", 1))
    return render(
        request,
        "customers/dataset_list.html",
        {
            "datasets": datasets,
        },
    )


@login_required
def dataset_detail(request, dataset_id):
    columns_queryset = DatasetColumn.objects.order_by("sheet_name", "order")
    dataset = get_object_or_404(
        get_accessible_datasets(request.user).prefetch_related(
            Prefetch("columns", queryset=columns_queryset),
        ).annotate(customer_count=Count("records")),
        pk=dataset_id,
    )
    columns_by_sheet = group_columns_by_sheet(dataset.columns.all())
    search_query = request.GET.get("q", "").strip()
    response_filter = request.GET.get("response", "").strip()
    records = get_filtered_records(dataset, search_query, response_filter)
    paginator = Paginator(records, 20)
    page_obj = paginator.get_page(request.GET.get("page"))
    records_by_sheet = build_records_by_sheet(page_obj.object_list, columns_by_sheet)
    response_stats = get_dataset_response_stats(dataset)
    import_summary = request.session.pop("last_import_summary", None)
    if import_summary and import_summary.get("dataset_id") != dataset.id:
        request.session["last_import_summary"] = import_summary
        import_summary = None
    active_filters = bool(search_query or response_filter)
    query_params = request.GET.copy()
    query_params.pop("page", None)

    return render(
        request,
        "customers/dataset_detail.html",
        {
            "dataset": dataset,
            "columns_by_sheet": columns_by_sheet,
            "records_by_sheet": records_by_sheet,
            "import_summary": import_summary,
            "page_obj": page_obj,
            "response_choices": CustomerRecord.RESPONSE_CHOICES,
            "response_filter": response_filter,
            "search_query": search_query,
            "active_filters": active_filters,
            "filtered_count": paginator.count,
            "query_string": query_params.urlencode(),
            "response_stats": response_stats,
        },
    )


@login_required
def dataset_export(request, dataset_id):
    columns_queryset = DatasetColumn.objects.order_by("sheet_name", "order")
    dataset = get_object_or_404(
        get_accessible_datasets(request.user).prefetch_related(
            Prefetch("columns", queryset=columns_queryset),
        ),
        pk=dataset_id,
    )
    search_query = request.GET.get("q", "").strip()
    response_filter = request.GET.get("response", "").strip()
    records = get_filtered_records(dataset, search_query, response_filter)
    workbook = build_dataset_export_workbook(dataset, records)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = f"{safe_export_filename(dataset.name)}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@login_required
def dataset_rename(request, dataset_id):
    dataset = get_object_or_404(get_accessible_datasets(request.user), pk=dataset_id)
    form = DatasetForm(request.POST or None, instance=dataset)

    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Dataset renamed successfully.")
        return redirect("customers:dataset_detail", dataset_id=dataset.id)

    return render(
        request,
        "customers/dataset_rename.html",
        {
            "dataset": dataset,
            "form": form,
        },
    )


@login_required
def dataset_delete(request, dataset_id):
    dataset = get_object_or_404(
        get_accessible_datasets(request.user).annotate(customer_count=Count("records")),
        pk=dataset_id,
    )

    if request.method == "POST":
        dataset_name = dataset.name
        dataset.delete()
        messages.success(request, f"Dataset '{dataset_name}' deleted successfully.")
        return redirect("customers:dataset_list")

    return render(
        request,
        "customers/dataset_confirm_delete.html",
        {
            "dataset": dataset,
        },
    )


@login_required
def customer_detail(request, dataset_id, record_id):
    dataset, record = get_dataset_record(request.user, dataset_id, record_id)
    columns = get_record_columns(dataset, record)

    return render(
        request,
        "customers/customer_detail.html",
        {
            "dataset": dataset,
            "record": record,
            "field_rows": build_field_rows(record, columns),
        },
    )


@login_required
def customer_edit(request, dataset_id, record_id):
    dataset, record = get_dataset_record(request.user, dataset_id, record_id)
    form = CustomerCRMForm(request.POST or None, instance=record)

    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Customer CRM information updated successfully.")
        return redirect(
            "customers:customer_detail",
            dataset_id=dataset.id,
            record_id=record.id,
        )

    return render(
        request,
        "customers/customer_edit.html",
        {
            "dataset": dataset,
            "record": record,
            "form": form,
            "field_rows": build_field_rows(record, get_record_columns(dataset, record)),
        },
    )


@login_required
def customer_delete(request, dataset_id, record_id):
    dataset, record = get_dataset_record(request.user, dataset_id, record_id)

    if request.method == "POST":
        record.delete()
        messages.success(request, "Customer record deleted successfully.")
        return redirect("customers:dataset_detail", dataset_id=dataset.id)

    return render(
        request,
        "customers/customer_confirm_delete.html",
        {
            "dataset": dataset,
            "record": record,
            "field_rows": build_field_rows(record, get_record_columns(dataset, record)),
        },
    )


@login_required
def import_excel(request):
    form = ExcelUploadForm(request.POST or None, request.FILES or None)
    inspection = None
    pending_import_token = None
    dataset_name = None

    if request.method == "POST" and request.POST.get("action") == "confirm_import":
        return confirm_import(request)

    if request.method == "POST" and form.is_valid():
        uploaded_file = form.cleaned_data["excel_file"]
        dataset_name = form.cleaned_data["dataset_name"]

        try:
            inspection = inspect_workbook(uploaded_file)
            pending_import_token = save_pending_upload(request, uploaded_file, dataset_name)
            messages.success(
                request,
                "Excel file inspected successfully. Review the structure, then confirm the import.",
            )
        except ImportError:
            messages.error(
                request,
                "Excel inspection requires pandas and openpyxl. Install the packages in requirements.txt and try again.",
            )
        except zipfile.BadZipFile:
            messages.error(request, "The uploaded file is not a readable .xlsx workbook.")
        except ValueError as exc:
            messages.error(request, f"Could not inspect this workbook: {exc}")
        except Exception:
            messages.error(
                request,
                "Something went wrong while reading the workbook. Please check that the file is not corrupted.",
            )

    return render(
        request,
        "customers/import_excel.html",
        {
            "form": form,
            "inspection": inspection,
            "dataset_name": dataset_name,
            "pending_import_token": pending_import_token,
        },
    )


def confirm_import(request):
    token = request.POST.get("pending_import_token", "")
    pending_upload = get_pending_upload(request, token)

    if not pending_upload:
        messages.error(
            request,
            "The pending import could not be found. Please upload the Excel file again.",
        )
        return redirect("customers:import_excel")

    try:
        with open(pending_upload["path"], "rb") as workbook_file:
            inspection = inspect_workbook(workbook_file, filename=pending_upload["filename"])
        dataset, import_summary = create_dataset_from_inspection(
            request.user,
            pending_upload["dataset_name"],
            inspection,
        )
        import_summary["dataset_id"] = dataset.id
        import_summary["dataset_name"] = dataset.name
        request.session["last_import_summary"] = import_summary
        delete_pending_upload(request, token)
        messages.success(request, "Import completed successfully.")
        return redirect("customers:dataset_detail", dataset_id=dataset.id)
    except ImportError:
        messages.error(
            request,
            "Excel import requires pandas and openpyxl. Install the packages in requirements.txt and try again.",
        )
    except FileNotFoundError:
        messages.error(request, "The temporary uploaded file was not found. Please upload it again.")
        delete_pending_upload(request, token)
    except zipfile.BadZipFile:
        messages.error(request, "The uploaded file is not a readable .xlsx workbook.")
        delete_pending_upload(request, token)
    except ValueError as exc:
        messages.error(request, f"Could not import this workbook: {exc}")
    except Exception:
        messages.error(
            request,
            "Something went wrong while importing the workbook. Please check that the file is not corrupted.",
        )

    return redirect("customers:import_excel")


def inspect_workbook(uploaded_file, filename=None):
    import pandas as pd
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    workbook_bytes = uploaded_file.read()
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)

    try:
        workbook = load_workbook(
            filename=io.BytesIO(workbook_bytes),
            read_only=True,
            data_only=True,
        )
    except InvalidFileException as exc:
        raise ValueError("openpyxl could not read this Excel file.") from exc

    sheets = []

    for worksheet in workbook.worksheets:
        max_row = worksheet.max_row or 0
        max_column = worksheet.max_column or 0
        header_values = next(
            worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
            (),
        )
        column_names = [
            format_column_name(value, index)
            for index, value in enumerate(header_values[:max_column], start=1)
        ]

        sheet_data = read_sheet_data(pd, workbook_bytes, worksheet.title)

        sheets.append(
            {
                "name": worksheet.title,
                "rows": len(sheet_data["rows"]),
                "columns": max_column,
                "column_names": column_names,
                "field_keys": build_field_keys(column_names),
                "data_rows": sheet_data["rows"],
                "preview_headers": column_names,
                "preview_rows": sheet_data["preview_rows"],
            }
        )

    workbook.close()

    return {
        "filename": filename or uploaded_file.name,
        "sheet_names": [sheet["name"] for sheet in sheets],
        "sheets": sheets,
    }


def read_sheet_data(pd, workbook_bytes, sheet_name):
    dataframe = pd.read_excel(
        io.BytesIO(workbook_bytes),
        sheet_name=sheet_name,
        engine="openpyxl",
        dtype=object,
    )

    dataframe = dataframe.where(pd.notnull(dataframe), "")
    data_rows = []

    for record in dataframe.to_dict(orient="records"):
        data_rows.append([serialize_cell_value(record.get(column, "")) for column in dataframe.columns])

    return {
        "rows": data_rows,
        "preview_rows": [
            [mask_sensitive_value(value) for value in row]
            for row in data_rows[:5]
        ],
    }


def create_dataset_from_inspection(owner, dataset_name, inspection):
    import_summary = {
        "filename": inspection["filename"],
        "rows_found": 0,
        "rows_imported": 0,
        "rows_skipped": 0,
        "errors": 0,
    }

    with transaction.atomic():
        dataset = Dataset.objects.create(
            owner=owner,
            name=dataset_name,
            original_filename=inspection["filename"],
        )

        for sheet in inspection["sheets"]:
            columns = []
            for index, column_name in enumerate(sheet["column_names"], start=1):
                columns.append(
                    DatasetColumn(
                        dataset=dataset,
                        sheet_name=sheet["name"],
                        original_name=column_name,
                        field_key=sheet["field_keys"][index - 1],
                        order=index,
                    )
            )
            DatasetColumn.objects.bulk_create(columns)

            records = []
            for row_offset, row_values in enumerate(sheet["data_rows"], start=2):
                import_summary["rows_found"] += 1
                if not row_has_value(row_values):
                    import_summary["rows_skipped"] += 1
                    continue

                original_data = {
                    field_key: row_values[index]
                    for index, field_key in enumerate(sheet["field_keys"])
                    if index < len(row_values)
                }
                records.append(
                    CustomerRecord(
                        dataset=dataset,
                        sheet_name=sheet["name"],
                        row_number=row_offset,
                        original_data=original_data,
                    )
                )
                import_summary["rows_imported"] += 1
            CustomerRecord.objects.bulk_create(records)

    return dataset, import_summary


def save_pending_upload(request, uploaded_file, dataset_name):
    pending_dir = Path(settings.BASE_DIR) / "tmp" / "excel_imports"
    pending_dir.mkdir(parents=True, exist_ok=True)

    token = uuid.uuid4().hex
    filename = Path(uploaded_file.name).name
    file_path = pending_dir / f"{token}.xlsx"

    with open(file_path, "wb") as destination:
        for chunk in uploaded_file.chunks():
            destination.write(chunk)

    pending_imports = request.session.get("pending_imports", {})
    pending_imports[token] = {
        "path": str(file_path),
        "filename": filename,
        "dataset_name": dataset_name,
    }
    request.session["pending_imports"] = pending_imports
    return token


def get_pending_upload(request, token):
    pending_imports = request.session.get("pending_imports", {})
    return pending_imports.get(token)


def delete_pending_upload(request, token):
    pending_imports = request.session.get("pending_imports", {})
    pending_upload = pending_imports.pop(token, None)

    if pending_upload:
        try:
            os.remove(pending_upload["path"])
        except FileNotFoundError:
            pass

    request.session["pending_imports"] = pending_imports


def build_field_keys(column_names):
    keys = []
    seen = {}

    for index, column_name in enumerate(column_names, start=1):
        base_key = slugify(column_name).replace("-", "_") or f"column_{index}"
        occurrence = seen.get(base_key, 0) + 1
        seen[base_key] = occurrence
        key = base_key if occurrence == 1 else f"{base_key}_{occurrence}"
        keys.append(key)

    return keys


def group_columns_by_sheet(columns):
    grouped = {}
    for column in columns:
        grouped.setdefault(column.sheet_name, []).append(column)
    return grouped


def build_records_by_sheet(records, columns_by_sheet, masked=True):
    grouped = []

    for sheet_name, columns in columns_by_sheet.items():
        sheet_records = []
        for record in records:
            if record.sheet_name != sheet_name:
                continue
            sheet_records.append(
                {
                    "record": record,
                    "values": [
                        format_record_value(record.original_data.get(column.field_key, ""), masked)
                        for column in columns
                    ],
                }
            )

        grouped.append(
            {
                "columns": columns,
                "records": sheet_records,
                "sheet_name": sheet_name,
            }
        )

    return grouped


def user_can_access_all_datasets(user):
    return user.is_staff or user.is_superuser


def get_admin_user_queryset():
    return get_user_model().objects.annotate(
        dataset_count=Count("datasets", distinct=True),
        customer_count=Count("datasets__records", distinct=True),
    )


def get_accessible_datasets(user):
    datasets = Dataset.objects.all()
    if user_can_access_all_datasets(user):
        return datasets
    return datasets.filter(owner=user)


def get_accessible_visits(user):
    visits = Visit.objects.all()
    if user_can_access_all_datasets(user):
        return visits
    return visits.filter(
        Q(customer__dataset__owner=user)
        | Q(manual_client__owner=user)
    )


def get_accessible_visit_or_404(user, visit_id):
    return get_object_or_404(
        get_accessible_visits(user).select_related(
            "customer",
            "customer__dataset",
            "manual_client",
            "manual_client__owner",
        ),
        pk=visit_id,
    )


def get_selected_dataset(request, datasets, source=None):
    source = source or request.GET
    dataset_id = source.get("dataset")
    if not dataset_id:
        return None
    try:
        return datasets.get(pk=dataset_id)
    except (Dataset.DoesNotExist, ValueError):
        return None


def get_dataset_record(user, dataset_id, record_id):
    dataset = get_object_or_404(get_accessible_datasets(user), pk=dataset_id)
    record = get_object_or_404(
        CustomerRecord,
        pk=record_id,
        dataset=dataset,
    )
    return dataset, record


def get_record_columns(dataset, record):
    return list(
        dataset.columns.filter(sheet_name=record.sheet_name).order_by("order")
    )


def build_field_rows(record, columns):
    return [
        {
            "label": column.original_name,
            "value": mask_sensitive_value(record.original_data.get(column.field_key, "")),
        }
        for column in columns
    ]


def build_manual_client_rows(client):
    return [
        {"label": "Client Name", "value": client.client_name},
        {"label": "Phone", "value": client.phone},
        {"label": "Email", "value": client.email},
        {"label": "Company", "value": client.company},
        {"label": "Location", "value": client.location},
        {"label": "Created", "value": client.created_at.strftime("%d %b %Y, %H:%M")},
        {"label": "Updated", "value": client.updated_at.strftime("%d %b %Y, %H:%M")},
    ]


def format_record_value(value, masked=True):
    if masked:
        return mask_sensitive_value(value)
    return value


def record_matches_search(record, search_query):
    query = search_query.casefold()
    values = [str(value) for value in record.original_data.values()]
    values.extend([record.get_response_display(), record.note])
    return any(query in value.casefold() for value in values)


def get_visit_customer_options(dataset, search_query=""):
    records = list(dataset.records.order_by("sheet_name", "row_number"))
    if not search_query:
        return records

    query = search_query.casefold()
    return [
        record for record in records
        if query in " ".join(str(value) for value in record.original_data.values()).casefold()
    ]


def get_filtered_records(dataset, search_query="", response_filter=""):
    records = dataset.records.order_by("sheet_name", "row_number")

    if response_filter:
        records = records.filter(response=response_filter)

    records = list(records)
    if search_query:
        records = [
            record for record in records
            if record_matches_search(record, search_query)
        ]

    return records


def get_dataset_response_stats(dataset):
    counts_by_response = {
        row["response"]: row["count"]
        for row in dataset.records.values("response").annotate(count=Count("id"))
    }
    return [
        {
            "value": value,
            "label": label,
            "count": counts_by_response.get(value, 0),
        }
        for value, label in CustomerRecord.RESPONSE_CHOICES
    ]


def build_dataset_export_workbook(dataset, records):
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)
    columns_by_sheet = group_columns_by_sheet(dataset.columns.order_by("sheet_name", "order"))
    records_by_sheet = build_records_by_sheet(records, columns_by_sheet, masked=False)

    if not records_by_sheet:
        workbook.create_sheet("Customers")

    for sheet in records_by_sheet:
        worksheet = workbook.create_sheet(safe_sheet_title(workbook, sheet["sheet_name"]))
        headers = [column.original_name for column in sheet["columns"]] + ["Response", "Note"]
        worksheet.append(headers)

        for item in sheet["records"]:
            record = item["record"]
            worksheet.append(item["values"] + [record.get_response_display(), record.note])

        for index, header in enumerate(headers, start=1):
            width = min(max(len(str(header)) + 2, 12), 40)
            worksheet.column_dimensions[get_column_letter(index)].width = width

    return workbook


def safe_export_filename(dataset_name):
    filename = slugify(dataset_name).replace("-", "_")
    return filename or "dataset_export"


def safe_sheet_title(workbook, sheet_name):
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", " ", sheet_name).strip() or "Customers"
    cleaned = cleaned[:31]
    candidate = cleaned
    suffix = 1

    while candidate in workbook.sheetnames:
        suffix += 1
        suffix_text = f" {suffix}"
        candidate = f"{cleaned[:31 - len(suffix_text)]}{suffix_text}"

    return candidate


def format_column_name(value, index):
    if value is None or str(value).strip() == "":
        return f"Unnamed column {index}"
    return str(value).strip()


def serialize_cell_value(value):
    if value is None or value == "":
        return ""

    if hasattr(value, "item"):
        value = value.item()

    if hasattr(value, "isoformat"):
        return value.isoformat()

    return value


def row_has_value(row_values):
    return any(value not in ("", None) for value in row_values)


def mask_sensitive_value(value):
    if value is None or value == "":
        return ""

    text = str(value)

    text = EMAIL_PATTERN.sub(mask_email_match, text)

    digit_count = len(DIGIT_PATTERN.findall(text))
    if digit_count >= 7:
        visible_digits = 0
        masked_characters = []
        for character in reversed(text):
            if character.isdigit():
                visible_digits += 1
                masked_characters.append(character if visible_digits <= 2 else "X")
            else:
                masked_characters.append(character)
        text = "".join(reversed(masked_characters))

    return text


def mask_email_match(match):
    local_part = match.group(1)
    domain = match.group(2)

    if len(local_part) <= 2:
        masked_local = "*" * len(local_part)
    else:
        masked_local = f"{local_part[0]}***{local_part[-1]}"

    return f"{masked_local}@{domain}"
